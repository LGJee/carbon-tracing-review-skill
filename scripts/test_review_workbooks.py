"""Local regression tests using synthetic workbooks, not biochemical evidence."""

import argparse
from contextlib import redirect_stdout
from io import StringIO
import json
import subprocess
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

sys.dont_write_bytecode = True
import review_workbooks as review

MAPPED = "[CH3:1][CH2:2][OH:3]>>[CH3:1][CH2:2][OH:3]"
SWAPPED = "[CH3:1][CH2:2][OH:3]>>[CH3:2][CH2:1][OH:3]"
PROJECT = None
MAPPING_SKILL = None


def fixture(path):
    workbook = Workbook()
    del workbook[workbook.active.title]
    equations = [
        "A[e] -> A[c]",
        "A[c] -> B[c]",
        "B[c] -> C[c]",
        "C[c] -> D[c]",
        "D[c] -> E[c]",
    ]
    ids = [0, "r1", "r2", "r3", "r4"]
    sheet = workbook.create_sheet("Nonzero")
    sheet.append(["ReactionID", "ReactionName", "Equation", "Flux"])
    for key, equation in zip(ids, equations):
        sheet.append([key, "Synthetic fixture", equation, 1])
    sheet = workbook.create_sheet("Metabolites")
    sheet.append(["MetaboliteID", "MetaboliteName", "SMILES_Model", "SMILES_Manual"])
    for i, name in enumerate(["A[e]", "A[c]", "B[c]", "C[c]", "D[c]", "E[c]"]):
        sheet.append(["m" + str(i), name, "CCO", None])
    for backend in review.BACKENDS:
        sheet = workbook.create_sheet(backend)
        sheet.append(
            [
                "ReactionID",
                "ReactionName",
                "Equation",
                "ReactionSMILES",
                "MappedReactionSMILES",
                "MappingDirection",
                "MappingConfident",
            ]
        )
        for key, equation in zip(ids, equations):
            sheet.append(
                [
                    key,
                    "Synthetic fixture",
                    equation,
                    "CCO>>CCO",
                    MAPPED,
                    "forward",
                    0.8273648273648273,
                ]
            )
    selection = workbook.create_sheet("Mapping_Selection")
    selection.append(
        [
            "ReactionID",
            "ReactionName",
            "Flux",
            "SelectedMapper",
            "ReviewRequired",
            "SelectionReason",
            "Equation",
            "MappingDirection",
            "SelectedMappedReactionSMILES",
        ]
    )
    final = workbook.create_sheet("Final_Mapping")
    final.append(
        [
            "ReactionID",
            "ReactionName",
            "Flux",
            "SelectedMapper",
            "ReviewRequired",
            "AIGenerated",
            "MappedReactionSMILES",
            "Equation",
            "MappingDirection",
        ]
    )
    for key, equation, flag in zip(
        ids, equations, [False, "TRUE", True, "FALSE", False]
    ):
        selection.append(
            [
                key,
                "Synthetic fixture",
                1,
                "LocalMapper",
                flag,
                "Fixture only",
                equation,
                "forward",
                None if key == "r2" else MAPPED,
            ]
        )
        if key != "r2":
            final.append(
                [
                    key,
                    "Synthetic fixture",
                    1,
                    "LocalMapper",
                    key in ("r1", "r3"),
                    False,
                    MAPPED,
                    equation,
                    "forward",
                ]
            )
    info = workbook.create_sheet("Mapping_Info")
    info.append(["Mapper", "SourceInputSHA256", "OptionsJSON"])
    info.append(["LocalMapper", "a" * 64, "{}"])
    info.sheet_state = "hidden"
    workbook.save(path)
    workbook.close()


def proposal(key, mapped=MAPPED):
    return dict(
        ReactionID=key,
        MappedReactionSMILES=mapped,
        scope="all",
        evidence_status="evidence-supported",
        rationale="Synthetic test fixture; not a scientific evidence claim.",
        limitations=["Fixture only"],
        evidence=[
            dict(
                url="https://example.org/fixture",
                accessed_on="2026-10-14",
                claim="Synthetic test metadata, not retrieved evidence.",
            )
        ],
    )


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="carbon-review-test-")
        self.folder = Path(self.temporary.name)
        self.source = self.folder / "Mixed_Culture_2_atom_mapping.xlsx"
        fixture(self.source)
        self.before = review.digest(self.source)
        # Fixtures must never enter the real project's scientific review archive.
        writer = review.archive_review

        def save_archive(report, targets, proposals, directory):
            if directory == PROJECT / "data/ai_mapping_reviews":
                directory = self.folder / "project/data/ai_mapping_reviews"
            return writer(report, targets, proposals, directory)

        patcher = patch.object(review, "archive_review", side_effect=save_archive)
        patcher.start()
        self.addCleanup(patcher.stop)
        index_writer = review.index_ai_mappings
        def save_index(report, targets, proposals, directory, species):
            if directory == PROJECT / "data":
                directory = self.folder / "project/data"
            return index_writer(report, targets, proposals, directory, species)
        index_patcher = patch.object(review, "index_ai_mappings", side_effect=save_index)
        index_patcher.start()
        self.addCleanup(index_patcher.stop)
        work_directory = review.ai_work_directory
        def isolated_work(project, fallback, stage):
            if project == PROJECT:
                project = self.folder / "isolated-project"
            return work_directory(project, fallback, stage)
        cache_patcher = patch.object(review, "ai_work_directory", side_effect=isolated_work)
        cache_patcher.start()
        self.addCleanup(cache_patcher.stop)

    def test_reduced_mapper_context_acceptance_rejection_and_complete_guard(self):
        import pandas as pd
        from copy import deepcopy
        from carbon_flow_tools.pipeline import finish_analysis, load_final_mappings
        from carbon_flow_tools.reactions import prepare_reactions
        root = self.folder / "reduced"
        root.mkdir()
        source = pd.DataFrame([["R", "Carrier-dependent", 1,
            "A[c] + ReducedProtein[c] => B[c] + OxidizedProtein[c]"]],
            columns=["ReactionID", "ReactionName", "Flux", "Equation"])
        metabolites = pd.DataFrame([["a", "A[c]", "CCO"], ["b", "B[c]", "CCO"],
            ["red", "ReducedProtein[c]", ""], ["ox", "OxidizedProtein[c]", ""]],
            columns=["MetaboliteID", "MetaboliteName", "SMILES_Model"])
        path = root / "input.xlsx"
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            source.to_excel(writer, sheet_name="Nonzero", index=False)
            metabolites.to_excel(writer, sheet_name="Metabolites", index=False)
        reactions = prepare_reactions(source, metabolites)
        if not hasattr(reactions[0], "omitted_components"):
            self.skipTest("Installed project predates reduced mapper inputs")
        reactions[0].mapped_smiles, reactions[0].status = MAPPED, "mapped"
        predictions = {name: deepcopy(reactions) for name in review.BACKENDS}
        finish_analysis("Reduced", dict(input_file=path, metabolite_sheet="Metabolites"),
            metabolites, predictions, [], list(predictions), root, stage="map")
        staged = root / "Reduced_2_atom_mapping.xlsx"
        wb = load_workbook(staged)
        try:
            targets, _ = review.review_targets(wb)
            self.assertEqual(len(targets["R"]["omitted_components"]), 2)
        finally:
            wb.close()
        for verdict in ("accepted", "rejected", "unresolved"):
            entry = dict(input_file=str(staged), input_sha256=review.digest(staged), replacements=[],
                reviews=[dict(ReactionID="R", AIReviewed=True, rule_verdict=verdict,
                    decision="unchanged" if verdict == "accepted" else "unresolved",
                    rationale="Synthetic biochemical-review fixture, not real evidence")])
            report = review.apply_workbook(entry, PROJECT, MAPPING_SKILL, root / verdict)
            loaded, _, _ = load_final_mappings(Path(report["output_file"]), {})
            self.assertEqual(loaded[0].status == "partial", verdict == "accepted")
        entry = dict(input_file=str(staged), input_sha256=review.digest(staged), replacements=[proposal("R")])
        report = review.apply_workbook(entry, PROJECT, MAPPING_SKILL, root / "false_complete")
        self.assertEqual(report["accepted"], 0)
        self.assertIn("cannot establish complete", report["reactions"][0]["reason"])

    def tearDown(self):
        self.assertEqual(review.digest(self.source), self.before)
        self.temporary.cleanup()

    def entry(self, replacements):
        return dict(
            input_file=str(self.source),
            input_sha256=self.before,
            replacements=replacements,
        )

    def apply(self, replacements, use_project=True, output="reviewed"):
        return review.apply_workbook(
            self.entry(replacements),
            PROJECT if use_project else None,
            MAPPING_SKILL,
            self.folder / output,
        )

    def test_flags_and_union_include_missing_final_rows(self):
        self.assertFalse(review.flag(" FALSE "))
        self.assertTrue(review.flag("true"))
        self.assertEqual(review.identity(0), "0")
        self.assertEqual(review.identity(12.0), "12")
        with self.assertRaises(ValueError):
            review.flag("Ture")
        with self.assertRaises(ValueError):
            review.flag(1)
        workbook = load_workbook(self.source)
        try:
            targets, issues = review.review_targets(workbook)
        finally:
            workbook.close()
        self.assertEqual(set(targets), {"r1", "r2", "r3"})
        self.assertTrue(targets["r2"]["missing_final_row"])
        self.assertFalse(issues)

    def test_blank_false_rows_are_reviewed_but_declared_inert_rows_are_exempt(self):
        workbook = load_workbook(self.source)
        try:
            selection, final = workbook["Mapping_Selection"], workbook["Final_Mapping"]
            a = {cell.value: cell.column for cell in selection[1]}
            b = {cell.value: cell.column for cell in final[1]}
            final.cell(1, final.max_column + 1, "MappingCoverage")
            coverage_column = final.max_column
            selection.cell(6, a["SelectedMappedReactionSMILES"], "")
            final.cell(5, b["MappedReactionSMILES"], "")
            targets, _ = review.review_targets(workbook)
            self.assertEqual(targets["r4"]["review_trigger"], "missing_mapping")
            for coverage in ("Boundary", "NoCarbon"):
                final.cell(5, coverage_column, coverage)
                targets, _ = review.review_targets(workbook)
                self.assertNotIn("r4", targets)
        finally:
            workbook.close()

    def test_blank_raw_smiles_can_use_exact_saved_project_inputs(self):
        original_fixture = self.source
        candidate = self.folder / "Blank_2_atom_mapping.xlsx"
        fixture(candidate)
        workbook = load_workbook(candidate)
        try:
            for name in review.BACKENDS:
                sheet = workbook[name]
                headers = {cell.value: cell.column for cell in sheet[1]}
                sheet.cell(3, headers["ReactionSMILES"], "")
                sheet.cell(3, headers["MappedReactionSMILES"], "")
            sheet = workbook["Final_Mapping"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            sheet.cell(3, headers["MappedReactionSMILES"], "")
            workbook.save(candidate)
        finally:
            workbook.close()
        report = review.apply_workbook(dict(input_file=str(candidate), input_sha256=review.digest(candidate),
            replacements=[proposal("r1")]), PROJECT, MAPPING_SKILL, self.folder / "blank_reviewed")
        self.assertEqual(report["accepted"], 1)

    def test_success_insertion_failure_and_provenance_preservation(self):
        report = self.apply(
            [
                proposal("r1", SWAPPED),
                proposal("r2"),
                proposal("r3", "[CH4:1]>>[CH4:1]"),
            ]
        )
        self.assertEqual(report["accepted"], 2)
        output = Path(report["output_file"])
        original, reviewed = load_workbook(self.source), load_workbook(output)
        try:
            for name in [
                "Nonzero",
                "Metabolites",
                "Mapping_Info",
                "Mapping_Selection",
            ] + list(review.BACKENDS):
                self.assertEqual(
                    list(original[name].values), list(reviewed[name].values)
                )
            self.assertEqual(reviewed["Mapping_Info"].sheet_state, "hidden")
            _, final = review.table(reviewed, "Final_Mapping", review.FINAL_COLUMNS)
            _, selection = review.table(
                reviewed, "Mapping_Selection", review.SELECTION_COLUMNS
            )
            self.assertEqual(len(final), 5)
            for key in ("r1", "r2"):
                self.assertIs(final[key]["AIGenerated"], True)
                self.assertIs(final[key]["AIReviewed"], True)
                self.assertEqual(final[key]["SelectedMapper"], "AI-assisted")
            self.assertEqual(final["r1"]["MappedReactionSMILES"], SWAPPED)
            self.assertEqual(selection["r1"]["SelectedMappedReactionSMILES"], MAPPED)
            self.assertIsNone(selection["r2"]["SelectedMappedReactionSMILES"])
            self.assertEqual(final["r2"]["MappedReactionSMILES"], MAPPED)
            self.assertIs(final["r1"]["ReviewRequired"], True)
            self.assertIs(final["r2"]["ReviewRequired"], True)
            self.assertIs(final["r3"]["AIGenerated"], False)
            self.assertIs(final["r3"]["AIReviewed"], False)
            self.assertEqual(final["r3"]["MappedReactionSMILES"], MAPPED)
            self.assertEqual(selection["r4"]["SelectedMappedReactionSMILES"], MAPPED)
        finally:
            original.close()
            reviewed.close()

    def test_hash_and_existing_completed_output_are_protected(self):
        wrong = self.entry([proposal("r1")])
        wrong["input_sha256"] = "0" * 64
        with self.assertRaisesRegex(ValueError, "SHA256"):
            review.apply_workbook(
                wrong, PROJECT, MAPPING_SKILL, self.folder / "reviewed"
            )
        report = self.apply([proposal("r1")])
        output = Path(report["output_file"])
        saved = review.digest(output)
        with self.assertRaisesRegex(ValueError, "retained"):
            self.apply([proposal("r1", SWAPPED)])
        self.assertEqual(review.digest(output), saved)

    def test_invalid_scope_evidence_map_and_unflagged_rows_are_unchanged(self):
        entries = [proposal("r1"), proposal("r2"), proposal("r3"), proposal("r4")]
        entries[0]["scope"] = "carbon"
        entries[1]["evidence_status"] = "unresolved"
        entries[2]["MappedReactionSMILES"] = (
            "[CH3:1][CH2:1][OH:3]>>[CH3:1][CH2:2][OH:3]"
        )
        report = self.apply(entries)
        self.assertEqual(report["accepted"], 0)
        self.assertIsNone(report["output_file"])
        self.assertTrue(all(row["status"] == "rejected" for row in report["reactions"]))

    def test_raw_input_mismatch_and_duplicate_proposals_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "Duplicated proposal"):
            self.apply([proposal("r1"), proposal("r1")])
        workbook = load_workbook(self.source)
        workbook["RXNMapper"]["D3"] = "C>>C"
        workbook.save(self.source)
        workbook.close()
        self.before = review.digest(self.source)
        report = self.apply([proposal("r1")])
        self.assertEqual(report["accepted"], 0)

    def test_no_project_still_runs_scientific_structural_validation(self):
        report = self.apply([proposal("r1")], use_project=False)
        self.assertEqual(report["accepted"], 1)
        self.assertEqual(report["project_validation"], "not_available")

    def test_invalid_evidence_and_cell_overflow_cannot_change_failed_rows(self):
        bad_evidence = proposal("r2")
        bad_evidence["evidence"][0]["url"] = 42
        bad_map = proposal("r1", "C" * 32768)
        report = self.apply([bad_evidence, bad_map])
        self.assertEqual(report["accepted"], 0)
        too_long = proposal("r3")
        too_long["rationale"] = "x" * 32768
        report = self.apply([proposal("r1"), too_long])
        self.assertEqual(report["accepted"], 2)
        reviewed = load_workbook(report["output_file"])
        original = load_workbook(self.source)
        try:
            _, before = review.table(
                original, "Mapping_Selection", review.SELECTION_COLUMNS
            )
            _, after = review.table(
                reviewed, "Mapping_Selection", review.SELECTION_COLUMNS
            )
            self.assertEqual(before["r3"], after["r3"])
        finally:
            reviewed.close()
            original.close()

    def test_numeric_zero_id_is_not_dropped(self):
        workbook = load_workbook(self.source)
        sheet = workbook["Mapping_Selection"]
        column = [c.value for c in sheet[1]].index("ReviewRequired") + 1
        sheet.cell(2, column, True)
        workbook.save(self.source)
        workbook.close()
        self.before = review.digest(self.source)
        report = self.apply([proposal(0)])
        self.assertEqual(report["accepted"], 1)
        reviewed = load_workbook(report["output_file"])
        try:
            _, rows = review.table(reviewed, "Final_Mapping", review.FINAL_COLUMNS)
            self.assertIs(rows["0"]["AIGenerated"], True)
        finally:
            reviewed.close()

    def test_inspect_command_uses_each_workbook_species(self):
        sources = [self.folder / (name + "_2_atom_mapping.xlsx") for name in
                   ("Escherichia_coli_non_IUPD", "Saccharomyces_cerevisiae_IUPD")]
        for source in sources:
            fixture(source)
        hashes = [review.digest(source) for source in sources]
        stream = StringIO()
        with patch.object(review, "cached_reviews", return_value=({}, [])) as cache:
            with redirect_stdout(stream):
                result = review.main(["inspect"] + [str(path) for path in sources]
                                     + ["--project", str(PROJECT)])
        self.assertEqual(result, 0)
        reports = json.loads(stream.getvalue())["workbooks"]
        self.assertEqual([row["input_file"] for row in reports],
                         [str(path.resolve()) for path in sources])
        self.assertEqual([row["input_sha256"] for row in reports], hashes)
        self.assertTrue(all(row["targets"] for row in reports))
        self.assertEqual([call[0][2] for call in cache.call_args_list],
                         ["Escherichia_coli", "Saccharomyces_cerevisiae"])
        completed = subprocess.run(
            [sys.executable, str(Path(review.__file__).resolve()), "inspect"]
            + [str(path) for path in sources] + ["--project", str(PROJECT)],
            cwd=str(self.folder), stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            universal_newlines=True)
        self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
        fresh = json.loads(completed.stdout)["workbooks"]
        self.assertEqual([row["model"] for row in fresh],
                         ["Escherichia_coli_non_IUPD", "Saccharomyces_cerevisiae_IUPD"])
        self.assertTrue(all(row["targets"] for row in fresh))
        self.assertEqual([review.digest(source) for source in sources], hashes)

    def test_multiple_workbooks_continue_after_one_failure(self):
        second = self.folder / "Other_Culture_2_atom_mapping.xlsx"
        fixture(second)
        invalid = self.entry([proposal("r1")])
        invalid["input_sha256"] = "0" * 64
        valid = dict(
            input_file=str(second),
            input_sha256=review.digest(second),
            replacements=[proposal("r1")],
        )
        proposals = self.folder / "proposals.json"
        proposals.write_text(
            json.dumps(dict(workbooks=[invalid, valid])), encoding="utf-8"
        )
        stream = StringIO()
        arguments = [
            "apply",
            str(proposals),
            "--project",
            str(PROJECT),
            "--mapping-skill",
            str(MAPPING_SKILL),
            "--output-dir",
            str(self.folder / "out"),
        ]
        with redirect_stdout(stream):
            result = review.main(arguments)
        report = json.loads(stream.getvalue())["workbooks"]
        self.assertEqual(result, 2)
        self.assertIn("error", report[0])
        self.assertEqual(report[1]["accepted"], 1)
        self.assertTrue(Path(report[1]["output_file"]).is_file())
        self.assertEqual(review.digest(second), valid["input_sha256"])
        delivery = json.loads(stream.getvalue())
        self.assertEqual(len(delivery["report_files"]), 2)
        self.assertEqual(len(list((self.folder / "out" / "Ai_reports").glob("*.md"))), 1)
        self.assertFalse(list((self.folder / "out").glob("*.json")))
        for files, workbook_report in zip(delivery["report_files"], report):
            self.assertEqual(files["input_file"], workbook_report["input_file"])
            self.assertTrue(Path(files["markdown"]).is_file())
            saved = json.loads(Path(files["json"]).read_text())
            self.assertEqual(saved["workbooks"], [workbook_report])

    def test_each_model_gets_one_report_combining_its_structure_and_mapping_changes(self):
        first = self.apply([proposal("r1")], output="first")
        second_source = self.folder / "Other_Culture_2_atom_mapping.xlsx"
        fixture(second_source)
        second = review.apply_workbook(dict(input_file=str(second_source),
            input_sha256=review.digest(second_source), replacements=[proposal("r1")]),
            PROJECT, MAPPING_SKILL, self.folder / "second")
        first["structure_changes"] = [dict(MetaboliteID="m_unique_A", MetaboliteName="Only A[c]",
                                          action="Added", SMILES="CCO")]
        second["structure_changes"] = [dict(MetaboliteID="m_unique_B", MetaboliteName="Only B[c]",
                                           action="Added", SMILES="CCO")]
        destination = self.folder / "per_model_reports"
        for own, other in ((first, second), (second, first)):
            files = review.write_delivery_report([own], destination)
            text = Path(files["markdown"]).read_text(encoding="utf-8")
            self.assertIn(own["structure_changes"][0]["MetaboliteName"], text)
            self.assertNotIn(other["structure_changes"][0]["MetaboliteName"], text)
            self.assertIn("| r1 | Synthetic fixture |", text)
            self.assertNotIn(Path(other["input_file"]).name, text)
            self.assertIn(review.model_name(Path(own["input_file"])), Path(files["markdown"]).name)
            data = json.loads(Path(files["json"]).read_text(encoding="utf-8"))
            self.assertEqual(data["workbooks"], [own])
        self.assertEqual(len(list((destination / "Ai_reports").glob("*.md"))), 2)
        self.assertFalse(list(destination.glob("*.md")))
        self.assertFalse(list(destination.glob("*.json")))
        with self.assertRaises(ValueError):
            review.write_delivery_report([first, second], destination)

    def test_ai_cache_is_disposable_and_retry_keeps_one_final_model_report(self):
        import shutil
        report = self.apply([proposal("r1")], output="delivery")
        destination = self.folder / "delivery"
        isolated_project = self.folder / "clean-project"
        data = isolated_project / "data"
        data.mkdir(parents=True)
        reusable = data / "ai_smiles.csv"
        reusable.write_text("Species,MetaboliteName,SMILES,Source\n", encoding="utf-8")
        files = review.write_delivery_report([report], destination, project=isolated_project)
        first_content = Path(files["markdown"]).read_bytes()
        second = review.write_delivery_report([report], destination, project=isolated_project)
        self.assertEqual(files["markdown"], second["markdown"])
        self.assertNotEqual(files["json"], second["json"])
        cache = isolated_project / ".cache/ai"
        self.assertIn(cache, Path(files["json"]).parents)
        self.assertEqual(Path(second["markdown"]).parent, destination / "Ai_reports")
        self.assertEqual(len(list((destination / "Ai_reports").glob("*.md"))), 1)
        self.assertTrue(list(cache.rglob("previous_report.md")))
        shutil.rmtree(str(cache))
        self.assertEqual(Path(second["markdown"]).read_bytes(), first_content)
        self.assertTrue(Path(report["output_file"]).is_file())
        self.assertTrue(reusable.is_file())

    def test_failed_retry_report_stays_in_cache_and_preserves_final_report(self):
        report = self.apply([proposal("r1")], output="delivery")
        destination = self.folder / "delivery"
        first = review.write_delivery_report([report], destination)
        before = Path(first["markdown"]).read_bytes()
        failed = dict(input_file=str(self.source), error="Synthetic retry failure")
        files = review.write_delivery_report([failed], destination)
        self.assertIn(destination / ".cache/ai", Path(files["markdown"]).parents)
        self.assertEqual(Path(first["markdown"]).read_bytes(), before)
        self.assertEqual(len(list((destination / "Ai_reports").glob("*.md"))), 1)

    def test_project_rejects_mismatched_equation_before_writing(self):
        workbook = load_workbook(self.source)
        for name in list(review.BACKENDS) + ["Mapping_Selection", "Final_Mapping"]:
            sheet = workbook[name]
            column = [c.value for c in sheet[1]].index("MappingDirection") + 1
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, 1).value == "r1":
                    sheet.cell(row, column, "reverse")
        workbook.save(self.source)
        workbook.close()
        self.before = review.digest(self.source)
        report = self.apply([proposal("r1")])
        self.assertEqual(report["accepted"], 0)

    def test_written_copy_is_consumed_by_step3_emu_and_versions_stay_distinct(self):
        import shutil
        # Add balanced external input/drain; the original fixture is an open chain.
        wb = load_workbook(self.source)
        ws = wb["Final_Mapping"]
        ws.cell(1, ws.max_column + 1, "MappingCoverage")
        for row in range(2, ws.max_row + 1):
            ws.cell(row, ws.max_column, "Complete")
        columns = [cell.value for cell in ws[1]]
        for key, equation in [("in", "-> A[e]"), ("out", "E[c] ->")]:
            wb["Nonzero"].append([key, "Boundary fixture", equation, 1])
            record = dict(ReactionID=key, ReactionName="Boundary fixture", Flux=1,
                          SelectedMapper="", ReviewRequired=False, AIGenerated=False,
                          MappedReactionSMILES="", Equation=equation,
                          MappingDirection="forward", MappingCoverage="Boundary")
            ws.append([record.get(field) for field in columns])
            selection = wb["Mapping_Selection"]
            record.update(SelectionReason="Boundary fixture", SelectedMappedReactionSMILES="")
            selection.append([record.get(cell.value) for cell in selection[1]])
        wb.save(self.source)
        wb.close()
        self.before = review.digest(self.source)
        report = self.apply([proposal("r1", SWAPPED), proposal("r2")])
        from carbon_flow_tools import pipeline
        from carbon_flow_tools.cli import main
        destination = self.folder / "reviewed"
        ordinary = destination / self.source.name
        shutil.copy2(str(self.source), str(ordinary))
        mapping_hashes = {str(path.resolve()): review.digest(path)
                          for path in [ordinary, Path(report["output_file"])]}
        options = dict(input_dir=str(destination), output_dir=str(destination),
                       carbon_sources={"Mixed_Culture": ["A[e]"]},
                       max_carbon_processes=1, carbon_memory_gb=None,
                       emu_allow_inorganic_carbon_recycling=True, emu_tolerance=1e-8,
                       min_available_memory_gb=.01, max_excel_size_mb=50,
                       tracing_python=sys.executable)
        with patch.object(pipeline, "run_mapper", side_effect=AssertionError("Mapper invoked")):
            with redirect_stdout(StringIO()):
                self.assertEqual(main("trace", options, []), 0)
        for suffix, expected in [("", 1), ("_AIGenerated", 0)]:
            output = destination / ("Mixed_Culture_3_carbon_flow" + suffix + ".xlsx")
            workbook = load_workbook(output)
            try:
                self.assertEqual(set(workbook.sheetnames), {"EMU_Metabolites", "EMU_BoundaryInputs"})
                records = list(workbook["EMU_Metabolites"].values)
                internal = {"FluxResidual", "Status", "StructureType", "KnownCarbonCount", "ContributionScope"}
                self.assertFalse(internal.intersection(records[0]))
                rows = [dict(zip(records[0], row)) for row in records[1:]]
                product = next(row for row in rows if row["MetaboliteName"] == "C[c]")
                self.assertAlmostEqual(product["Unknown"], expected)
                self.assertAlmostEqual(product["A[e]"], 1 - expected)
            finally:
                workbook.close()
        manifest = json.loads((destination / "carbon_flow_run.json").read_text())
        self.assertEqual(manifest["method"], "EMU")
        self.assertEqual(len(manifest["models"]), 2)
        for status in manifest["models"]:
            self.assertTrue(status["emu_completed"])
            self.assertEqual(status["stop_reason"], "emu_solved")
            self.assertEqual(status["mapping_sha256"], mapping_hashes[status["mapping_file"]])
            self.assertNotIn("completed_rounds", status)
        self.assertFalse((destination / "carbon_flow_summary.xlsx").exists())
        self.assertTrue(Path(report["output_file"]).is_file())

    def test_archive_retains_inputs_evidence_and_failed_proposals(self):
        report = self.apply(
            [proposal("r1"), proposal("r2", "[CH4:1]>>[CH4:1]")], use_project=False
        )
        path = Path(report["archive_file"])
        self.assertEqual(path.parent, self.folder / "reviewed/data/ai_mapping_reviews")
        archive = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(archive["input_sha256"], self.before)
        records = {row["target"]["ReactionID"]: row for row in archive["reactions"]}
        self.assertEqual(records["r1"]["target"]["ReactionSMILES"], "CCO>>CCO")
        self.assertEqual(
            records["r1"]["proposal"]["evidence"], proposal("r1")["evidence"]
        )
        self.assertEqual(
            records["r1"]["outcome"]["checks"][0]["status"], "structural_checks_passed"
        )
        self.assertEqual(records["r2"]["outcome"]["status"], "rejected")
        book = load_workbook(self.source)
        try:
            targets, _ = review.review_targets(book)
        finally:
            book.close()
        candidates, issues = review.cached_reviews(self.folder / "reviewed", targets)
        self.assertFalse(issues)
        self.assertEqual(set(candidates), {"r1"})
        self.assertTrue(candidates["r1"][0]["requires_revalidation"])
        self.assertEqual(
            candidates["r1"][0]["proposal"]["MappedReactionSMILES"], MAPPED
        )
        cached = candidates["r1"][0]["proposal"]
        cached["MappedReactionSMILES"] = "[CH4:1]>>[CH4:1]"
        with self.assertRaises(ValueError):
            review.check_proposal(
                cached,
                targets["r1"],
                review.scientific_validator(MAPPING_SKILL),
                {},
                None,
            )
        # Removing CSV rows disables old cache entries without deleting their evidence.
        index = self.folder / "reviewed/data/ai_atom_mappings.csv"
        previous = index.read_bytes()
        index.write_text("Species,ReactionID,MappedReactionSMILES,Source\n", encoding="utf-8-sig")
        self.assertEqual(review.cached_reviews(self.folder / "reviewed", targets)[0], {})
        self.assertTrue(Path(report["archive_file"]).is_file())
        index.unlink()
        self.assertEqual(review.cached_reviews(self.folder / "reviewed", targets)[0], {})
        index.write_bytes(previous)
        targets["r1"]["MappingDirection"] = "reverse"
        self.assertEqual(
            review.cached_reviews(self.folder / "reviewed", targets)[0], {}
        )
        targets["r1"]["MappingDirection"] = "forward"
        targets["r1"]["input_variants"] = ["CC[O-]>>CC[O-]"]
        self.assertEqual(
            review.cached_reviews(self.folder / "reviewed", targets)[0], {}
        )

    def test_rejected_archive_is_not_reused_and_bad_json_is_reported(self):
        report = self.apply([proposal("r1", "[CH4:1]>>[CH4:1]")], use_project=False)
        self.assertIsNone(report["output_file"])
        self.assertTrue(Path(report["archive_file"]).is_file())
        book = load_workbook(self.source)
        try:
            targets, _ = review.review_targets(book)
        finally:
            book.close()
        directory = self.folder / "reviewed/data/ai_mapping_reviews"
        (directory / "broken.json").write_text("{", encoding="utf-8")
        candidates, issues = review.cached_reviews(self.folder / "reviewed", targets)
        self.assertEqual(candidates, {})
        self.assertEqual(len(issues), 1)

    def test_archive_write_failure_reports_but_retains_delivered_workbook(self):
        with patch.object(
            review, "archive_review", side_effect=PermissionError("read-only data")
        ):
            report = self.apply([proposal("r1")], use_project=False)
        self.assertTrue(Path(report["output_file"]).is_file())
        self.assertIn("read-only data", report["archive_error"])

    def test_existing_review_archives_are_preserved(self):
        first = self.apply([proposal("r1")], output="reviewed")
        first_path = Path(first["archive_file"])
        before = review.digest(first_path)
        second = self.apply([proposal("r1", SWAPPED)], output="second")
        self.assertNotEqual(first["archive_file"], second["archive_file"])
        self.assertEqual(first_path.parent, Path(second["archive_file"]).parent)
        self.assertEqual(review.digest(first_path), before)

    def test_review_only_distinguishes_retained_unresolved_and_screening(self):
        entry = self.entry([])
        entry["reviews"] = [
            dict(
                ReactionID="r1",
                decision="unchanged",
                AIReviewed=True,
                rationale="Scientific review retained the existing map.",
            ),
            dict(
                ReactionID="r3",
                decision="unresolved",
                AIReviewed=True,
                rationale="Scientific analysis did not resolve the mechanism.",
            ),
            dict(
                ReactionID="r2",
                decision="unresolved",
                AIReviewed=False,
                rationale="Input screening only; scientific review not performed.",
            ),
        ]
        report = review.apply_workbook(
            entry, PROJECT, MAPPING_SKILL, self.folder / "out"
        )
        self.assertEqual(report["accepted"], 0)
        self.assertEqual(report["ai_reviewed"], 2)
        self.assertEqual(report["project_validation"], "not_run")
        workbook = load_workbook(report["output_file"])
        try:
            columns, rows = review.table(
                workbook, "Final_Mapping", review.FINAL_COLUMNS
            )
            self.assertEqual(columns[columns.index("ReviewRequired") + 1], "AIReviewed")
            self.assertNotIn("r2", rows)
            for key in ("r1", "r3"):
                self.assertIs(rows[key]["AIReviewed"], True)
                self.assertIs(rows[key]["AIGenerated"], False)
                self.assertEqual(rows[key]["MappedReactionSMILES"], MAPPED)
        finally:
            workbook.close()
        paths = review.write_delivery_report([report], self.folder / "out")
        text = Path(paths["markdown"]).read_text(encoding="utf-8")
        self.assertIn("| Total | 2 | 0 | 1 | 1 | incomplete |", text)
        self.assertNotIn("### r2", text)
        self.assertEqual(report["review_counts"]["screened_only"], 1)
        self.assertEqual(report["remaining_review_ids"], ["r2"])
        self.assertEqual(report["review_status"], "incomplete")
        self.assertEqual(
            len(json.loads(Path(paths["json"]).read_text())["workbooks"][0]["reviews"]),
            3,
        )
        archive = json.loads(Path(report["archive_file"]).read_text())
        self.assertEqual(archive["reviews"], entry["reviews"])

    def test_true_flags_are_separate_from_missing_only_targets(self):
        workbook = load_workbook(self.source)
        try:
            selection, final = workbook["Mapping_Selection"], workbook["Final_Mapping"]
            a = {cell.value: cell.column for cell in selection[1]}
            b = {cell.value: cell.column for cell in final[1]}
            selection.cell(4, a["ReviewRequired"], False)  # Missing final row, no TRUE flag.
            selection.cell(6, a["SelectedMappedReactionSMILES"], "")
            final.cell(5, b["MappedReactionSMILES"], "")
            targets, issues = review.review_targets(workbook)
            self.assertFalse(issues)
            counts = review.target_counts(targets)
            self.assertEqual(counts, dict(review_required_selection=1, review_required_final=2,
                review_required_union=2, missing_mapping_only=2, targets=4))
            self.assertEqual(set(targets), {"r1", "r2", "r3", "r4"})
            self.assertNotIn("0", targets)
        finally:
            workbook.close()

    def test_screening_all_targets_is_incomplete_and_cli_returns_two(self):
        entry = self.entry([])
        entry["reviews"] = [dict(ReactionID=key, decision="unresolved", AIReviewed=False,
            rationale="Input screening only.") for key in ("r1", "r2", "r3")]
        # A previous TRUE flag is provenance, not scientific analysis in this run.
        workbook = load_workbook(self.source)
        review.mark_reviewed(workbook, {"r1"})
        workbook.save(self.source)
        workbook.close()
        entry["input_sha256"] = review.digest(self.source)
        self.before = entry["input_sha256"]
        proposals = self.folder / "screening.json"
        proposals.write_text(json.dumps(dict(workbooks=[entry])), encoding="utf-8")
        stream = StringIO()
        with redirect_stdout(stream):
            result = review.main(["apply", str(proposals), "--project", str(PROJECT),
                "--mapping-skill", str(MAPPING_SKILL), "--output-dir", str(self.folder / "out")])
        self.assertEqual(result, 2)
        delivery = json.loads(stream.getvalue())
        self.assertEqual(delivery["review_status"], "incomplete")
        report = delivery["workbooks"][0]
        self.assertEqual(report["review_counts"]["scientifically_reviewed"], 0)
        self.assertEqual(report["review_counts"]["remaining"], 3)
        self.assertEqual(set(report["remaining_review_ids"]), {"r1", "r2", "r3"})
        text = Path(delivery["report_files"][0]["markdown"]).read_text(encoding="utf-8")
        self.assertNotIn("### r", text)
        self.assertNotIn("Input screening only.", text)
        self.assertIn("review is not complete; 3 reactions", text)

    def test_complete_reviews_and_no_targets_have_success_exit_status(self):
        entry = self.entry([])
        entry["reviews"] = [dict(ReactionID=key, decision="unresolved", AIReviewed=True,
            rationale="Scientific analysis found insufficient evidence for a replacement.")
            for key in ("r1", "r2", "r3")]
        for no_targets in (False, True):
            if no_targets:
                workbook = load_workbook(self.source)
                for sheet, field in (("Mapping_Selection", "SelectedMappedReactionSMILES"),
                                     ("Final_Mapping", "MappedReactionSMILES")):
                    columns = {cell.value: cell.column for cell in workbook[sheet][1]}
                    for cells in workbook[sheet].iter_rows(min_row=2):
                        workbook[sheet].cell(cells[0].row, columns["ReviewRequired"], False)
                        workbook[sheet].cell(cells[0].row, columns[field], MAPPED)
                # Restore the deliberately missing final row from the fixture.
                row = list(next(workbook["Final_Mapping"].iter_rows(min_row=2, max_row=2, values_only=True)))
                row[0] = "r2"
                row[7] = "B[c] -> C[c]"
                workbook["Final_Mapping"].append(row)
                workbook.save(self.source)
                workbook.close()
                entry["reviews"] = []
                entry["input_sha256"] = review.digest(self.source)
                self.before = entry["input_sha256"]
            proposals = self.folder / ("none.json" if no_targets else "complete.json")
            proposals.write_text(json.dumps(dict(workbooks=[entry])), encoding="utf-8")
            stream = StringIO()
            with redirect_stdout(stream):
                result = review.main(["apply", str(proposals), "--project", str(PROJECT),
                    "--mapping-skill", str(MAPPING_SKILL), "--output-dir",
                    str(self.folder / ("none" if no_targets else "complete"))])
            self.assertEqual(result, 0, stream.getvalue())
            report = json.loads(stream.getvalue())["workbooks"][0]
            self.assertEqual(report["review_status"], "no_targets" if no_targets else "complete")
            self.assertEqual(report["remaining_review_ids"], [])
            self.assertEqual(report["review_counts"]["unresolved"], 0 if no_targets else 3)

    def test_readable_report_lists_only_applied_changes_and_keeps_audit_in_json(self):
        entry = self.entry([proposal("r1")])
        entry["reviews"] = [dict(ReactionID="r3", decision="unchanged", AIReviewed=True,
                                rationale="Detailed scientific analysis retained this map.")]
        report = review.apply_workbook(entry, PROJECT, MAPPING_SKILL, self.folder / "out")
        report["structure_changes"] = [dict(MetaboliteID="test_met", MetaboliteName="A|name[c]",
                                              action="Added", SMILES="CCO")]
        paths = review.write_delivery_report([report], self.folder / "Ai_reports")
        text = Path(paths["markdown"]).read_text(encoding="utf-8")
        self.assertIn("test_met | A\\|name[c] | Added | CCO", text)
        self.assertIn("| r1 | Synthetic fixture | A[c] -> B[c] |", text)
        self.assertNotIn("| r3 |", text)
        self.assertNotIn("Detailed scientific analysis retained this map.", text)
        self.assertNotIn("Input SHA256:", text)
        self.assertNotIn("Carbon-rule verdict:", text)
        saved = json.loads(Path(paths["json"]).read_text(encoding="utf-8"))["workbooks"][0]
        self.assertEqual(saved["reviews"], entry["reviews"])
        accepted = next(row for row in saved["reactions"] if row["status"] == "accepted")
        self.assertEqual(accepted["MappedReactionSMILES"], MAPPED)
        self.assertTrue(accepted["evidence"])

    def test_structure_changes_require_matching_workbook_and_immutable_archive(self):
        import hashlib
        archive_data = dict(parent_workbook_sha256="a" * 64, structures=[dict(
            identity=dict(MetaboliteID="m1", MetaboliteName="A[c]"),
            proposal=dict(SMILES="CCO"))])
        content = json.dumps(archive_data).encode("utf-8")
        archive = self.folder / (hashlib.sha256(content).hexdigest() + ".json")
        archive.write_bytes(content)
        metadata = dict(output_sha256=self.before, input_sha256="a" * 64,
                        structure_archive=str(archive), completed_structures=1)
        sidecar = self.source.with_suffix(".structures.json")
        sidecar.write_text(json.dumps(metadata), encoding="utf-8")
        changes = review.completed_structure_changes(self.source, self.before)
        self.assertEqual(changes, [dict(MetaboliteID="m1", MetaboliteName="A[c]",
                                        action="Added", SMILES="CCO")])
        with self.assertRaises(ValueError):
            review.completed_structure_changes(self.source, "0" * 64)
        archive.write_bytes(content + b" ")
        with self.assertRaises(ValueError):
            review.completed_structure_changes(self.source, self.before)

    def test_bad_review_metadata_cannot_mark_a_row(self):
        for changes in (
            {"AIReviewed": "TRUE"},
            {"decision": "correct"},
            {"ReactionID": "r4"},
            {"rationale": ""},
        ):
            entry = self.entry([])
            row = dict(
                ReactionID="r1",
                decision="unresolved",
                AIReviewed=True,
                rationale="Incomplete evidence after analysis.",
            )
            row.update(changes)
            entry["reviews"] = [row]
            with self.assertRaises(ValueError):
                review.apply_workbook(
                    entry, PROJECT, MAPPING_SKILL, self.folder / "bad"
                )
        self.assertFalse((self.folder / "bad").exists())

    def test_existing_review_column_is_moved_without_losing_flags(self):
        workbook = load_workbook(self.source)
        try:
            sheet = workbook["Final_Mapping"]
            column = sheet.max_column + 1
            sheet.cell(1, column, "AIReviewed")
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, column, row == 2)
            review.mark_reviewed(workbook, {"r3"})
            headers, records = review.table(
                workbook, "Final_Mapping", review.FINAL_COLUMNS
            )
            self.assertEqual(headers.count("AIReviewed"), 1)
            self.assertEqual(
                headers.index("AIReviewed"), headers.index("ReviewRequired") + 1
            )
            self.assertIs(records["0"]["AIReviewed"], True)
            self.assertIs(records["r3"]["AIReviewed"], True)
            self.assertIs(records["r1"]["AIReviewed"], False)
        finally:
            workbook.close()

    def test_report_separates_unassessed_and_validation_rejections(self):
        report = self.apply([proposal("r1", "[CH4:1]>>[CH4:1]"), proposal("r4")])
        paths = review.write_delivery_report([report], self.folder / "Ai_reports")
        text = Path(paths["markdown"]).read_text(encoding="utf-8")
        self.assertIn("| Total | 0 | 0 | 0 | 0 | incomplete |", text)
        self.assertNotIn("### r1", text)
        self.assertEqual(report["review_counts"]["validation_rejected"], 1)
        self.assertEqual(set(report["remaining_review_ids"]), {"r1", "r2", "r3"})
        second = review.write_delivery_report([report], self.folder / "Ai_reports")
        self.assertNotEqual(paths, second)
        self.assertTrue(Path(paths["markdown"]).is_file())


class AIRecordTests(unittest.TestCase):
    def test_success_only_reverse_direction_and_provenance_merge(self):
        import csv
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            original = directory / "manual_atom_mappings.csv"
            original.write_text("Species,ReactionID,MappedReactionSMILES\n")
            before = original.read_bytes()
            accepted = proposal("R", "[CH3:1][CH2:2]O>>[CH3:1][CH:2]=O")
            rejected = proposal("bad")
            targets = {"R": dict(MappingDirection="reverse"), "bad": dict(MappingDirection="forward")}
            report = dict(archive_file="ai_mapping_reviews/test.json", reactions=[
                dict(ReactionID="R", status="accepted"), dict(ReactionID="bad", status="rejected")])
            for _ in range(2):
                review.index_ai_mappings(report, targets, [accepted, rejected], directory, "Escherichia_coli")
            with (directory / "ai_atom_mappings.csv").open(encoding="utf-8-sig") as stream:
                reader = csv.DictReader(stream)
                self.assertEqual(reader.fieldnames, ["Species", "ReactionID", "MappedReactionSMILES", "Source"])
                rows = list(reader)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["ReactionID"], "R")
            self.assertEqual(rows[0]["MappedReactionSMILES"], ">>".join(reversed(accepted["MappedReactionSMILES"].split(">>"))))
            self.assertIn("archive: ai_mapping_reviews/test.json", rows[0]["Source"])
            self.assertEqual(original.read_bytes(), before)


class CarbonRuleReviewTests(unittest.TestCase):
    def test_real_partial_handoff_and_scientific_verdicts(self):
        sys.path.insert(0, str(PROJECT))
        from scripts.test_carbon_mapping_rules import fixture as rule_fixture
        from carbon_flow_tools.pipeline import finish_analysis, load_final_mappings
        import pandas as pd
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, metabolites, original, _, _, _ = rule_fixture()
            input_file = root / "input.xlsx"
            with pd.ExcelWriter(input_file, engine="openpyxl") as writer:
                source.to_excel(writer, sheet_name="Reactions", index=False)
                metabolites.to_excel(writer, sheet_name="Metabolites", index=False)
            finish_analysis("Rules", dict(input_file=input_file, metabolite_sheet="Metabolites"),
                            metabolites, {"LocalMapper": original}, [], ["LocalMapper"], root, stage="map")
            path = root / "Rules_2_atom_mapping.xlsx"
            wb = load_workbook(path)
            try:
                targets, _ = review.review_targets(wb)
                self.assertEqual(len(targets["B"]["partial_mappings"]), 2)
                self.assertIsNone(targets["B"]["ReactionSMILES"])
                self.assertNotIn("No usable unmapped ReactionSMILES", targets["B"]["issues"])
                original_sheets = {name: list(wb[name].values) for name in wb.sheetnames if name != "Final_Mapping"}
            finally:
                wb.close()
            for verdict in ("accepted", "rejected", "unresolved"):
                entry = dict(input_file=str(path), input_sha256=review.digest(path), replacements=[],
                             reviews=[dict(ReactionID="B", AIReviewed=True, rule_verdict=verdict,
                                           decision="unchanged" if verdict == "accepted" else "unresolved",
                                           rationale="Synthetic review for behavior testing, not scientific evidence.")])
                with patch.object(review, "archive_review", return_value="test-only"):
                    report = review.apply_workbook(entry, PROJECT, MAPPING_SKILL, root / verdict)
                checked = load_workbook(report["output_file"])
                try:
                    _, rows = review.table(checked, "Final_Mapping", review.FINAL_COLUMNS)
                    self.assertTrue(rows["B"]["ReviewRequired"])
                    self.assertTrue(rows["B"]["AIReviewed"])
                    self.assertFalse(rows["B"]["AIGenerated"])
                    self.assertEqual(rows["B"]["MappingCoverage"], "Partial" if verdict == "accepted" else "Unresolved")
                    for name, values in original_sheets.items():
                        self.assertEqual(list(checked[name].values), values)
                finally:
                    checked.close()
                loaded, _, _ = load_final_mappings(report["output_file"], {})
                self.assertEqual(bool(loaded[0].partial_maps), verdict == "accepted")
            entry["reviews"][0].pop("rule_verdict")
            with self.assertRaises(ValueError):
                review.apply_workbook(entry, PROJECT, MAPPING_SKILL, root / "missing_verdict")




class AIRecordTableTests(unittest.TestCase):
    def test_readable_order_encoding_and_deleted_rows_stay_deleted(self):
        import csv
        fields = ["Species", "MetaboliteName", "SMILES", "Source"]
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            entries = [dict(Species="Escherichia_coli", MetaboliteName=name,
                            SMILES="C", Source="https://example.org/test") for name in ("Z", "A")]
            path = review.save_ai_records(root, "ai_smiles.csv", fields, entries)
            self.assertTrue(path.read_bytes().startswith(b"\xef\xbb\xbf"))
            with path.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.DictReader(stream))
            self.assertEqual([row["MetaboliteName"] for row in rows], ["A", "Z"])
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=fields)
                writer.writeheader()
                writer.writerow(rows[1])
            review.save_ai_records(root, path.name, fields, [dict(entries[0], MetaboliteName="M")])
            rows = review.read_ai_records(root, path.name, fields)
            self.assertEqual([row["MetaboliteName"] for row in rows], ["M", "Z"])

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project", type=Path, required=True)
    parser.add_argument("--mapping-skill", type=Path, required=True)
    args, other = parser.parse_known_args()
    PROJECT = args.project.resolve()
    MAPPING_SKILL = args.mapping_skill.resolve()
    unittest.main(argv=[sys.argv[0]] + other)
