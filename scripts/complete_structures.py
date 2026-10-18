"""Persist evidence-backed missing structures and remap only affected reactions."""

import argparse
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
import csv
from datetime import datetime
import importlib.util
import json
from pathlib import Path
import re
import sys

import review_workbooks as review


def context(path, project):
    api = review.project_api(project)
    if api is None:
        raise ValueError("Structure completion requires a compatible carbon-tracing project")
    import pandas as pd
    with pd.ExcelFile(path, engine="openpyxl") as excel:
        sheets = api["input_sheets"](excel)
    reactions, source, metabolites = api["pipeline"].read_input(path, *sheets)
    return api, sheets, reactions, source, metabolites


def identity_context(model, row, species=None):
    from carbon_flow_tools.reactions import clean
    from carbon_flow_tools.smiles_completion import species_name
    return dict(Model=model, Species=species or species_name(model), **{field: clean(getattr(row, field, "")) for field in
                ("MetaboliteID", "MetaboliteName", "MolecularFormula", "Charge", "InChIKey", "InChI", "PubChemCID")})


def workbook_species(path, project):
    return review.review_species(path, project)


def inspect(path, project):
    from carbon_flow_tools.reactions import get_smiles
    _, _, _, _, metabolites = context(path, project)
    missing = []
    species = workbook_species(path, project)
    directory = project / "data/ai_structure_reviews"
    archives, issues = [], []
    for item in sorted(directory.glob("*.json")) if directory.is_dir() else []:
        try:
            if review.digest(item) != item.stem:
                raise ValueError("Archive hash mismatch")
            records = json.loads(item.read_text(encoding="utf-8"))["structures"]
            if not isinstance(records, list) or any(not isinstance(row, dict) or
                    not isinstance(row.get("identity"), dict) or not isinstance(row.get("proposal"), dict)
                    for row in records):
                raise ValueError("Invalid structure archive records")
            archives.extend(records)
        except (ValueError, KeyError, TypeError, OSError) as error:
            issues.append(dict(path=str(item), error=str(error)))
    for row in metabolites.itertuples(index=False):
        if not get_smiles(row):
            identity = identity_context(review.model_name(path), row, species)
            missing.append(dict(identity, cached_candidates=[entry for entry in archives
                                 if entry.get("identity") == identity]))
    return dict(input_file=str(path), input_sha256=review.digest(path),
                missing_structures=missing, cache_issues=issues)


def validate_structure(row, proposal):
    from rdkit import Chem
    from rdkit.Chem import rdMolDescriptors
    from carbon_flow_tools.carbon_mapping_rules import generic_pool
    from carbon_flow_tools.reactions import clean, get_smiles
    if get_smiles(row):
        raise ValueError("Only missing structures may be completed: " + row.MetaboliteID)
    if proposal.get("MetaboliteName") != row.MetaboliteName:
        raise ValueError("Metabolite ID/name mismatch")
    if generic_pool(row.MetaboliteName):
        raise ValueError("An unspecified macromolecular pool has no unique molecular structure")
    text = proposal.get("SMILES", "")
    molecule = Chem.MolFromSmiles(text) if isinstance(text, str) and text else None
    if (molecule is None or not molecule.GetNumAtoms() or
            any(atom.GetAtomicNum() == 0 for atom in molecule.GetAtoms())):
        raise ValueError("A valid fully specified molecular structure is required")
    formula = clean(getattr(row, "MolecularFormula", ""))
    observed = rdMolDescriptors.CalcMolFormula(molecule)
    # Model formulas often store charge separately from the elemental formula.
    uncharged = lambda value: re.sub(r"[+-]\d*$", "", value)
    if formula and uncharged(formula) != uncharged(observed):
        raise ValueError("Structure formula differs from the model: {} versus {}".format(observed, formula))
    charge = clean(getattr(row, "Charge", ""))
    if charge and float(charge) != Chem.GetFormalCharge(molecule):
        raise ValueError("Structure charge differs from the model")
    key = clean(getattr(row, "InChIKey", ""))
    if key and Chem.MolToInchiKey(molecule) != key:
        raise ValueError("Structure InChIKey differs from the model")
    if any(not isinstance(proposal.get(field), str) or not proposal[field].strip()
           for field in ("database_id", "rationale")):
        raise ValueError("Record the exact database identity and identity rationale")
    evidence = proposal.get("evidence")
    if not isinstance(evidence, list) or not evidence:
        raise ValueError("Structure evidence is required")
    for item in evidence:
        if not isinstance(item, dict) or not str(item.get("url", "")).startswith(("https://", "http://")):
            raise ValueError("Each evidence record requires a source URL")
        datetime.strptime(item.get("accessed_on", ""), "%Y-%m-%d")
        if not isinstance(item.get("supported_claim"), str) or not item["supported_claim"].strip():
            raise ValueError("Each source requires its supported identity claim")
    return Chem.MolToSmiles(molecule, isomericSmiles=True)


def persist(project, model, accepted, parent_hash, manual_file=None):
    """Save AI discoveries separately from manual records, retaining immutable evidence."""
    from carbon_flow_tools.smiles_completion import load_manual, species_name, metabolite_name
    from rdkit import Chem
    manual = Path(manual_file) if manual_file is not None else project / "data/manual_smiles.csv"
    ai_file = manual.with_name("ai_smiles.csv")
    rows = (load_manual(manual) if manual.is_file() else []) + (
        load_manual(ai_file, with_source=True) if ai_file.is_file() else [])
    for entry in accepted:
        identity, proposal = entry["identity"], entry["proposal"]
        species = identity.get("Species") or species_name(model)
        name = metabolite_name(identity["MetaboliteName"])
        matched = [row for row in rows if row["Species"] == species
                   and metabolite_name(row["MetaboliteName"]) == name]
        if any(Chem.MolToSmiles(Chem.MolFromSmiles(row["SMILES"])) != proposal["SMILES"] for row in matched):
            raise ValueError("Conflicting manual/AI structure; resolve explicitly before completing structures")
        rows.append(dict(Species=species, MetaboliteName=name, SMILES=proposal["SMILES"]))
    archive = dict(version=1, parent_workbook_sha256=parent_hash, structures=accepted)
    content = json.dumps(archive, indent=2, sort_keys=True) + "\n"
    import hashlib
    directory = project / "data/ai_structure_reviews"
    directory.mkdir(parents=True, exist_ok=True)
    destination = directory / (hashlib.sha256(content.encode("utf-8")).hexdigest() + ".json")
    if not destination.exists():
        destination.write_text(content, encoding="utf-8")
    additions = []
    for entry in accepted:
        identity, proposal = entry["identity"], entry["proposal"]
        sources = [item["url"] for item in proposal["evidence"]]
        sources.append("archive: ai_structure_reviews/" + destination.name)
        additions.append(dict(Species=identity.get("Species") or species_name(model),
            MetaboliteName=metabolite_name(identity["MetaboliteName"]), SMILES=proposal["SMILES"],
            Source=" | ".join(dict.fromkeys(sources))))
    review.save_ai_records(ai_file.parent, ai_file.name,
                          ["Species", "MetaboliteName", "SMILES", "Source"], additions)
    return destination


def update_rows(workbook, sheet_name, records, key="ReactionID"):
    """Patch keyed rows while preserving unrelated values and columns."""
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    for record in records:
        for field in record:
            if field not in headers:
                headers.append(field)
                sheet.cell(1, len(headers), field)
    positions = {review.identity(row[headers.index(key)].value): row[0].row
                 for row in sheet.iter_rows(min_row=2) if row[headers.index(key)].value is not None}
    for record in records:
        row = positions.get(review.identity(record[key]), sheet.max_row + 1)
        for field, value in record.items():
            cell = sheet.cell(row, headers.index(field) + 1)
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"


def complete(entry, project, output_dir):
    import pandas as pd
    from openpyxl import load_workbook
    from carbon_flow_tools import config, mapper_runner
    from carbon_flow_tools.cli import build_mapper_configs
    from carbon_flow_tools.consensus import select_mappings
    from carbon_flow_tools.carbon_mapping_rules import apply_rules, apply_manual_mappings, carbon_structure_issues
    from carbon_flow_tools.reactions import prepare_reactions, parse_side, split_equation
    path = Path(entry["input_file"]).resolve()
    before = review.digest(path)
    if entry.get("input_sha256") != before:
        raise ValueError("Input SHA256 changed since structure inspection")
    model = review.model_name(path)
    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / path.name
    if output.exists() or output.resolve() == path:
        raise ValueError("Use an unused staging directory; never overwrite the source or a prior delivery")
    api, sheets, original, source, metabolites = context(path, project)
    with pd.ExcelFile(path, engine="openpyxl") as excel:
        required = set(review.BACKENDS) | {"Mapping_Selection", "Final_Mapping", "Mapping_Info"}
        if not required <= set(excel.sheet_names):
            raise ValueError("Incremental completion requires all three mapper sheets and mapping provenance")
    proposals = entry.get("structures", [])
    if not isinstance(proposals, list) or not proposals:
        raise ValueError("Provide at least one evidence-backed missing structure")
    rows = {row.MetaboliteID: row for row in metabolites.itertuples(index=False)}
    accepted, seen = [], set()
    species = workbook_species(path, project)
    for proposal in proposals:
        key = proposal["MetaboliteID"]
        if key in seen or key not in rows:
            raise ValueError("Unknown or duplicated metabolite ID: " + str(key))
        seen.add(key)
        accepted.append(dict(identity=identity_context(model, rows[key], species),
                             proposal=dict(proposal, SMILES=validate_structure(rows[key], proposal))))
    specification = importlib.util.spec_from_file_location("structure_review_step2", str(project / "step2_atom_mapping.py"))
    step2 = importlib.util.module_from_spec(specification)
    specification.loader.exec_module(step2)
    settings = step2.settings()
    # Store validated discoveries even if missing environments subsequently block remapping.
    manual_file = (project / settings["structure_options"]["manual_file"]).resolve()
    archive = persist(project, model, accepted, before, manual_file)
    for field in ("SMILES_Retrieved", "StructureSource", "StructureReference", "StructureStatus"):
        if field not in metabolites:
            metabolites[field] = ""
        metabolites[field] = metabolites[field].astype(object)
    if "AIStructureCompleted" not in metabolites:
        metabolites["AIStructureCompleted"] = False
    if "StructureReviewRequired" not in metabolites:
        metabolites["StructureReviewRequired"] = False
    names = set()
    for item in accepted:
        key, proposal = item["identity"]["MetaboliteID"], item["proposal"]
        mask = metabolites.MetaboliteID.eq(key)
        names.add(item["identity"]["MetaboliteName"])
        for field, value in dict(SMILES_Retrieved=proposal["SMILES"], AIStructureCompleted=True,
                                 StructureSource="AI-assisted structure lookup", StructureReference=str(archive),
                                 StructureReviewRequired=True,
                                 StructureStatus="Identity researched; reaction mapping still requires review").items():
            metabolites.loc[mask, field] = value
    config.configure(settings["batch_size"], settings["max_reaction_atoms"])
    prepared = prepare_reactions(source, metabolites)
    affected = [reaction for reaction in prepared if names.intersection(
                name for side in split_equation(reaction.equation) for _, name in parse_side(side))]
    if not affected:
        raise ValueError("Completed metabolites do not participate in the saved reaction network")
    predictions = {backend: deepcopy(affected) for backend in review.BACKENDS}
    pending = [reaction.reaction_id for reaction in affected if reaction.status == "ready"]
    if pending:
        configs = build_mapper_configs(settings)
        for backend in review.BACKENDS:
            mapper_runner.validate_mapper_config(backend, configs[backend])
        workers, threads = config.parallel_capacity(3, settings["max_mapper_processes"], settings["mapper_memory_gb"])
        with ThreadPoolExecutor(max_workers=workers) as executor:
            tasks = [executor.submit(mapper_runner.run_mapper, model, predictions[backend],
                                     backend, configs[backend], threads) for backend in review.BACKENDS]
            for task in tasks:
                task.result()
    selected, choices = select_mappings(predictions, carbon_structure_issues(metabolites))
    apply_manual_mappings(species, selected, choices, metabolites,
                         project / settings.get("manual_mapping_file", "data/manual_atom_mappings.csv"))
    partials = apply_rules(selected, choices, metabolites)
    choices["ReviewRequired"] = True
    choices["SelectionReason"] += "; input structures completed with AI assistance; biochemical review required"
    workbook = load_workbook(path)
    try:
        changed = metabolites.loc[metabolites.MetaboliteID.isin(seen)].astype(object)
        update_rows(workbook, sheets[1], changed.where(changed.notna(), None).to_dict("records"), key="MetaboliteID")
        for backend in review.BACKENDS:
            if backend not in workbook.sheetnames:
                raise ValueError("All three mapper sheets are required for incremental completion")
            update_rows(workbook, backend, api["pipeline"].prediction_table(predictions[backend]).to_dict("records"))
        fields = [cell.value for cell in workbook["Mapping_Selection"][1]]
        update_rows(workbook, "Mapping_Selection", choices[[field for field in fields if field in choices]].to_dict("records"))
        final = api["pipeline"].final_mapping_table(selected, choices)
        if "AIReviewed" in [cell.value for cell in workbook["Final_Mapping"][1]]:
            final["AIReviewed"] = False
        update_rows(workbook, "Final_Mapping", final.to_dict("records"))
        ids = {reaction.reaction_id for reaction in affected}
        if "Partial_Mapping" in workbook.sheetnames:
            values = list(workbook["Partial_Mapping"].values)
            old = pd.DataFrame(values[1:], columns=values[0])
            partials = pd.concat([old.loc[~old.ReactionID.astype(str).isin(ids)], partials], ignore_index=True)
        if not partials.empty:
            api["pipeline"].write_sheet(workbook, "Partial_Mapping", partials)
        elif "Partial_Mapping" in workbook.sheetnames:
            del workbook["Partial_Mapping"]
        values = list(workbook["Mapping_Info"].values)
        metadata = pd.DataFrame(values[1:], columns=values[0])
        if "ReactionID" in metadata:
            metadata = metadata.loc[~metadata.ReactionID.astype(str).isin(ids)]
        details = [dict(Mapper=backend, ReactionID=r.reaction_id, MappingStatus=r.status,
                        MappingTemplate=r.template, MappingInputEquation=r.mapping_equation,
                        OmittedComponentsJSON=json.dumps(r.omitted_components, sort_keys=True),
                        StructureParentWorkbookSHA256=before,
                        StructureReviewArchive=str(archive))
                   for backend, reactions in predictions.items() for r in reactions]
        api["pipeline"].write_sheet(workbook, "Mapping_Info", pd.concat([metadata, pd.DataFrame(details)], ignore_index=True))
        api["pipeline"].save_output(workbook, output)
    finally:
        workbook.close()
    try:
        api["pipeline"].load_final_mappings(output, {})
        if review.digest(path) != before:
            raise ValueError("Source workbook changed during completion")
    except Exception:
        output.unlink()
        raise
    report = dict(input_file=str(path), input_sha256=before, output_file=str(output),
                  output_sha256=review.digest(output), structure_archive=str(archive),
                  manual_file=str(manual_file), completed_structures=len(accepted),
                  affected_reactions=sorted(ids), mapper_reactions=pending,
                  note="Mapper-generated atom maps are not AI-generated atom maps; reaction review is still pending")
    output.with_suffix(".structures.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("inspect", "complete"))
    parser.add_argument("input", type=Path)
    parser.add_argument("--project", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    project = review.find_project(args.project, [args.input])
    review.project_api(project)
    if args.command == "inspect":
        result = inspect(args.input.resolve(), project)
    else:
        if args.output_dir is None:
            parser.error("complete requires --output-dir pointing to an unused staging directory")
        entry = json.loads(args.input.read_text(encoding="utf-8"))
        result = complete(entry, project, args.output_dir.resolve())
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()