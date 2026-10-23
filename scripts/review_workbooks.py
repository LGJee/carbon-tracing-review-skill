"""Extract flagged reactions and apply validated review proposals to new copies."""

import argparse
import csv
from copy import copy, deepcopy
from datetime import datetime
import hashlib
import importlib
import importlib.util
import json
import math
from pathlib import Path
import re
import sys
import tempfile
from urllib.parse import urlparse
from zipfile import BadZipFile

from openpyxl import load_workbook

sys.dont_write_bytecode = True
BACKENDS = ("LocalMapper", "RXNMapper", "EnzymeMap")
FINAL_COLUMNS = {
    "ReactionID",
    "ReactionName",
    "Flux",
    "Equation",
    "MappingDirection",
    "SelectedMapper",
    "ReviewRequired",
    "AIGenerated",
    "MappedReactionSMILES",
}
SELECTION_COLUMNS = (FINAL_COLUMNS - {"AIGenerated", "MappedReactionSMILES"}) | {
    "SelectedMappedReactionSMILES",
    "SelectionReason",
}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def identity(value):
    if value is None or isinstance(value, bool):
        raise ValueError("ReactionID must be nonempty, not Boolean")
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("ReactionID must be finite")
        if value.is_integer():
            value = int(value)
    value = str(value).strip()
    if not value:
        raise ValueError("ReactionID must be nonempty")
    return value


def flag(value):
    if type(value) is bool:
        return value
    if isinstance(value, str) and value.strip().lower() in ("true", "false"):
        return value.strip().lower() == "true"
    raise ValueError("Unrecognized ReviewRequired flag: {!r}".format(value))


def table(workbook, name, required):
    if name not in workbook.sheetnames:
        raise ValueError("Missing sheet: " + name)
    sheet = workbook[name]
    headers = [cell.value for cell in sheet[1]]
    while headers and headers[-1] is None:
        headers.pop()
    if len(set(headers)) != len(headers) or not required <= set(headers):
        raise ValueError("Duplicated or missing columns in " + name)
    rows = {}
    for cells in sheet.iter_rows(min_row=2, max_col=len(headers)):
        values = [cell.value for cell in cells]
        if all(value is None for value in values):
            continue
        row = dict(zip(headers, values))
        key = identity(row["ReactionID"])
        if key in rows:
            raise ValueError("Duplicated ReactionID in {}: {}".format(name, key))
        row["_row"] = cells[0].row
        rows[key] = row
    return headers, rows


def review_targets(workbook):
    _, selection = table(workbook, "Mapping_Selection", SELECTION_COLUMNS)
    _, final = table(workbook, "Final_Mapping", FINAL_COLUMNS)
    if not set(final) <= set(selection):
        raise ValueError("Final_Mapping contains IDs absent from Mapping_Selection")
    if "Mapping_Info" not in workbook.sheetnames:
        raise ValueError("Missing Mapping_Info; preserve step-2 provenance")
    raw = {}
    for name in BACKENDS:
        if name in workbook.sheetnames:
            _, raw[name] = table(
                workbook,
                name,
                {
                    "ReactionID",
                    "Equation",
                    "ReactionSMILES",
                    "MappingDirection",
                },
            )
    if not raw:
        raise ValueError("No raw mapper sheet supplies ReactionSMILES")
    partials = {}
    for name in ("Mapping_Info", "Partial_Mapping"):
        if name not in workbook.sheetnames:
            continue
        rows = list(workbook[name].values)
        headers = rows[0] if rows else ()
        records = [dict(zip(headers, values)) for values in rows[1:] if any(v is not None for v in values)]
        if name == "Mapping_Info":
            records = [{key: value for key, value in row.items() if key != "RecordType" and value is not None}
                       for row in records if row.get("RecordType") == "PartialMapping"]
        if not records:
            continue
        if not {"ReactionID", "Rule", "CarbonAtomMap", "SourceSMILES", "ProductSMILES"} <= set(headers):
            raise ValueError("Saved partial correspondences are missing rule context")
        for record in records:
            partials.setdefault(identity(record["ReactionID"]), []).append(record)
        break
    omissions = {}
    info = list(workbook["Mapping_Info"].values)
    if info and {"ReactionID", "OmittedComponentsJSON"} <= set(info[0]):
        for values in info[1:]:
            record = dict(zip(info[0], values))
            text = record.get("OmittedComponentsJSON")
            if record.get("ReactionID") is not None and text:
                items = json.loads(text)
                if not isinstance(items, list) or any(not isinstance(item, dict) for item in items):
                    raise ValueError("Invalid saved mapper omission context")
                key = identity(record["ReactionID"])
                if key in omissions and omissions[key] != items:
                    raise ValueError("Inconsistent mapper omission context: " + key)
                omissions[key] = items
    metabolite_context = []
    for sheet in workbook:
        values = iter(sheet.values)
        headers = next(values, ())
        if {"MetaboliteID", "MetaboliteName"} <= set(headers):
            fields = {"MetaboliteID", "MetaboliteName", "MolecularFormula", "Charge",
                      "SMILES_Manual", "SMILES_Model", "SMILES_Yeast9", "SMILES_Retrieved",
                      "StructureSource", "StructureReference", "StructureStatus", "StructureReviewRequired"}
            metabolite_context.extend({k: v for k, v in zip(headers, row) if k in fields}
                                      for row in values if any(value is not None for value in row))
    targets, issues = {}, []
    for key, row in selection.items():
        errors, needs_review = [], False
        for name, record in [
            ("Mapping_Selection", row),
            ("Final_Mapping", final.get(key)),
        ]:
            if record is not None:
                try:
                    needs_review |= flag(record["ReviewRequired"])
                except ValueError as error:
                    errors.append(name + ": " + str(error))
        if errors:
            issues.append({"ReactionID": key, "errors": errors[:]})
        current = final.get(key)
        coverage = str((current or {}).get("MappingCoverage", "")).strip()
        exempt = coverage in ("Boundary", "NoCarbon")
        blank_mapping = (current is None or not str(current.get("MappedReactionSMILES") or "").strip()
                         or not str(row.get("SelectedMappedReactionSMILES") or "").strip())
        if not needs_review and not (blank_mapping and not exempt):
            continue
        variants, predictions = [], []
        for name, records in raw.items():
            candidate = records.get(key)
            if candidate is None:
                continue
            predictions.append(
                dict(Mapper=name, **{k: v for k, v in candidate.items() if k != "_row"})
            )
            for field in ("Equation", "MappingDirection"):
                if str(candidate[field]).strip() != str(row[field]).strip():
                    errors.append(name + " has inconsistent " + field)
            text = candidate["ReactionSMILES"]
            if isinstance(text, str) and text.strip() and text not in variants:
                variants.append(text)
        if not variants and not partials.get(key):
            errors.append("No usable unmapped ReactionSMILES")
        if key in final:
            for field in ("Equation", "MappingDirection"):
                if str(final[key][field]).strip() != str(row[field]).strip():
                    errors.append("Final_Mapping has inconsistent " + field)
        targets[key] = {
            "ReactionID": key,
            "ReactionName": row["ReactionName"],
            "Equation": row["Equation"],
            "MappingDirection": row["MappingDirection"],
            "Flux": row["Flux"],
            "ReactionSMILES": variants[0] if variants else None,
            "input_variants": variants,
            "predictions": predictions,
            "selection": {k: v for k, v in row.items() if k != "_row"},
            "final": {k: v for k, v in final.get(key, {}).items() if k != "_row"},
            "partial_mappings": partials.get(key, []),
            "omitted_components": omissions.get(key, []),
            "metabolites": [item for item in metabolite_context
                            if str(item.get("MetaboliteName", "")) in str(row["Equation"])],
            "missing_final_row": key not in final,
            "review_trigger": "flagged" if needs_review else "missing_mapping",
            "issues": errors,
        }
    return targets, issues


def target_counts(targets):
    """Count worksheet flags separately from the union used as the review queue."""
    def is_true(record):
        try:
            return bool(record) and flag(record["ReviewRequired"])
        except ValueError:
            return False  # Unknown flags are reported by review_targets, never guessed.
    selection = {key for key, row in targets.items() if is_true(row["selection"])}
    final = {key for key, row in targets.items() if is_true(row["final"])}
    return dict(review_required_selection=len(selection),
                review_required_final=len(final),
                review_required_union=len(selection | final),
                missing_mapping_only=len(set(targets) - selection - final),
                targets=len(targets))


def review_progress(report):
    """Screening and rejected proposals alone never complete scientific review."""
    decisions = {identity(row["ReactionID"]): row for row in report.get("reviews", [])}
    counts = dict(report["target_counts"], scientifically_reviewed=0, replaced=0,
                  retained=0, unresolved=0, screened_only=0, unassessed=0,
                  validation_rejected=0)
    remaining = []
    for outcome in report["reactions"]:
        if not outcome["flagged"]:
            continue
        key = identity(outcome["ReactionID"])
        decision = decisions.get(key, {})
        accepted = outcome["status"] == "accepted"
        analyzed = accepted or decision.get("AIReviewed") is True
        counts["scientifically_reviewed"] += int(analyzed)
        counts["validation_rejected"] += int(outcome["status"] == "rejected")
        if accepted:
            counts["replaced"] += 1
        elif analyzed:
            counts["retained" if decision["decision"] == "unchanged" else "unresolved"] += 1
        else:
            remaining.append(key)
            counts["screened_only" if decision else "unassessed"] += 1
    counts["remaining"] = len(remaining)
    if (counts["scientifically_reviewed"] != report["ai_reviewed"]
            or counts["replaced"] != report["accepted"]):
        raise ValueError("Review counts differ from applied scientific decisions")
    if counts["targets"] != counts["scientifically_reviewed"] + counts["remaining"]:
        raise ValueError("Review target/progress counts do not reconcile")
    if counts["targets"] != sum(counts[key] for key in
            ("replaced", "retained", "unresolved", "screened_only", "unassessed")):
        raise ValueError("Review outcome counts do not reconcile")
    status = "no_targets" if not counts["targets"] else "incomplete" if remaining else "complete"
    return dict(review_counts=counts, remaining_review_ids=remaining, review_status=status)


def model_name(path):
    match = re.fullmatch(r"(.+)_2_atom_mapping\.xlsx", path.name)
    if not match or path.name.startswith("~$"):
        raise ValueError("Use the original <model>_2_atom_mapping.xlsx filename")
    return match.group(1)


def find_project(explicit, paths):
    candidates = (
        [Path(explicit).resolve()]
        if explicit
        else [Path.cwd()]
        + [parent for path in paths for parent in Path(path).resolve().parents]
    )
    for directory in candidates:
        if (directory / "carbon_flow_tools/consensus.py").is_file():
            return directory
    if explicit:
        raise ValueError("The specified project lacks carbon_flow_tools/consensus.py")
    return None


def project_api(project):
    if project is None:
        return None
    sys.path.insert(0, str(project))
    pipeline = importlib.import_module("carbon_flow_tools.pipeline")
    excel = importlib.import_module("carbon_flow_tools.excel_io")
    consensus = importlib.import_module("carbon_flow_tools.consensus")
    if Path(pipeline.__file__).resolve().parent.parent != project:
        raise ValueError("Another project is already imported; use a fresh interpreter")
    required = [
        (pipeline, "read_input"),
        (pipeline, "save_output"),
        (excel, "input_sheets"),
        (consensus, "mapping_signatures"),
    ]
    if any(not callable(getattr(module, name, None)) for module, name in required):
        raise ValueError("Incompatible project review interfaces")
    return dict(
        pipeline=pipeline,
        input_sheets=excel.input_sheets,
        mapping_signatures=consensus.mapping_signatures,
    )


def scientific_validator(mapping_skill):
    directory = (
        Path(mapping_skill).resolve()
        if mapping_skill
        else (Path(__file__).resolve().parents[2] / "enzymatic-atom-mapping")
    )
    script = directory / "scripts/validate_mapping.py"
    if not script.is_file():
        raise ValueError(
            "Scientific validator missing; locate enzymatic-atom-mapping with --mapping-skill"
        )
    spec = importlib.util.spec_from_file_location(
        "scientific_mapping_validator", str(script)
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.validate_mapping


def check_proposal(proposal, target, validator, prepared, api):
    if target.get("omitted_components") or getattr(prepared.get(target["ReactionID"]), "omitted_components", []):
        raise ValueError("A reduced mapper input cannot establish complete original-reaction coverage; review it as partial or complete the missing structures first")
    if target["issues"]:
        raise ValueError("; ".join(target["issues"]))
    if getattr(prepared.get(target["ReactionID"]), "structure_error", ""):
        raise ValueError(prepared[target["ReactionID"]].structure_error)
    if proposal.get("scope", "all") != "all":
        raise ValueError("Partial scope cannot replace a default all-atom mapping")
    status = proposal.get("evidence_status")
    if status not in ("evidence-supported", "partially-supported"):
        raise ValueError("Evidence remains unresolved or unrecognized")
    if (
        not isinstance(proposal.get("rationale"), str)
        or not proposal["rationale"].strip()
    ):
        raise ValueError("A scientific rationale is required")
    evidence = proposal.get("evidence")
    if not isinstance(evidence, list) or not evidence:
        raise ValueError("Actual evidence references are required")
    for item in evidence:
        if not isinstance(item, dict):
            raise ValueError("Evidence references must be objects")
        if any(
            not isinstance(item.get(field), str) or not item[field].strip()
            for field in ("url", "claim", "accessed_on")
        ):
            raise ValueError(
                "Evidence URL, claim and access date must be nonempty text"
            )
        url = urlparse(item.get("url", ""))
        if (
            url.scheme not in ("http", "https")
            or not url.netloc
            or not item.get("claim")
        ):
            raise ValueError("Evidence needs a source URL and supported claim")
        datetime.strptime(item.get("accessed_on", ""), "%Y-%m-%d")
    marked = proposal.get("MappedReactionSMILES")
    if not isinstance(marked, str) or not marked.strip():
        raise ValueError("A mapped reaction string is required")
    if len(marked) > 32767:
        raise ValueError(
            "Mapped reaction exceeds Excel's cell limit; no truncated map will be saved"
        )
    reports = [
        validator(original, marked, scope="all")
        for original in target["input_variants"]
    ]
    if api is not None:
        key = target["ReactionID"]
        if key not in prepared:
            raise ValueError("ReactionID is absent from the prepared project input")
        reaction = deepcopy(prepared[key])
        if (
            reaction.equation != target["Equation"]
            or ("forward" if reaction.flux >= 0 else "reverse")
            != target["MappingDirection"]
            or not math.isclose(
                float(reaction.flux),
                float(target["Flux"]),
                rel_tol=1e-12,
                abs_tol=1e-15,
            )
        ):
            raise ValueError(
                "Project reaction input/direction does not match review target"
            )
        reaction.mapped_smiles = marked
        api["mapping_signatures"](reaction)
    return reports


def update_row(workbook, key, target, proposal):
    row = target["selection"]
    final = workbook["Final_Mapping"]
    columns, existing = table(workbook, "Final_Mapping", FINAL_COLUMNS)
    if key in existing:
        number = existing[key]["_row"]
    else:
        number = final.max_row + 1
        for index, column in enumerate(columns, 1):
            if column in row:
                final.cell(number, index, row[column])
            if number > 2:
                final.cell(number, index)._style = copy(
                    final.cell(number - 1, index)._style
                )
        final.cell(number, columns.index("ReactionID") + 1, key)
    for column, value in {
        "MappedReactionSMILES": proposal["MappedReactionSMILES"],
        "AIGenerated": True,
        "SelectedMapper": "AI-assisted",
    }.items():
        final.cell(number, columns.index(column) + 1, value)
    if "MappingCoverage" in columns:
        final.cell(number, columns.index("MappingCoverage") + 1, "Complete")


def apply_rule_verdict(workbook, key, target, decision):
    """Disable unsupported rules without changing the original audit sheet."""
    reduced_partial = target.get("omitted_components") and target.get("final", {}).get("MappingCoverage") == "Partial"
    if not target.get("partial_mappings") and not reduced_partial:
        if "rule_verdict" in decision:
            raise ValueError("rule_verdict requires saved partial-rule or mapper-omission context")
        return
    if not decision["AIReviewed"]:
        if "rule_verdict" in decision:
            raise ValueError("A rule verdict requires scientific review")
        return
    verdict = decision.get("rule_verdict")
    if verdict not in ("accepted", "rejected", "unresolved"):
        raise ValueError("Scientific rule reviews require an explicit rule_verdict")
    if decision["decision"] != ("unchanged" if verdict == "accepted" else "unresolved"):
        raise ValueError("Rule verdict and review decision disagree")
    columns, rows = table(workbook, "Final_Mapping", FINAL_COLUMNS)
    if key not in rows or "MappingCoverage" not in columns:
        raise ValueError("Rule review requires the project's MappingCoverage contract")
    if verdict != "accepted":
        sheet = workbook["Final_Mapping"]
        sheet.cell(rows[key]["_row"], columns.index("MappingCoverage") + 1, "Unresolved")
        sheet.cell(rows[key]["_row"], columns.index("MappedReactionSMILES") + 1).value = None


def mark_reviewed(workbook, reviewed):
    sheet = workbook["Final_Mapping"]
    columns, rows = table(workbook, "Final_Mapping", FINAL_COLUMNS)
    position = columns.index("ReviewRequired") + 2
    if "AIReviewed" in columns:
        old = columns.index("AIReviewed") + 1
        values = {key: flag(row["AIReviewed"]) for key, row in rows.items()}
        sheet.delete_cols(old)
        columns.remove("AIReviewed")
        position = columns.index("ReviewRequired") + 2
    else:
        values = {}
    sheet.insert_cols(position)
    sheet.cell(1, position, "AIReviewed")
    for number in range(1, sheet.max_row + 1):
        sheet.cell(number, position)._style = copy(
            sheet.cell(number, position - 1)._style
        )
    for key, row in rows.items():
        sheet.cell(row["_row"], position, key in reviewed or values.get(key, False))


def same_sheet_data(expected, actual):
    if expected.sheetnames != actual.sheetnames:
        return False
    for name in expected.sheetnames:
        a, b = expected[name], actual[name]
        if (a.max_row, a.max_column, a.sheet_state) != (
            b.max_row,
            b.max_column,
            b.sheet_state,
        ):
            return False
        if list(a.values) != list(b.values):
            return False
    return True


def output_directory(explicit, project, source):
    if explicit:
        return Path(explicit).resolve()
    plan = project / "data/analysis_plan.json" if project else None
    if plan is not None and plan.is_file():
        settings = json.loads(plan.read_text(encoding="utf-8"))
        directory = settings.get("output_dir")
        if not isinstance(directory, str) or not directory:
            raise ValueError("Analysis plan has no valid output_dir")
        return (project / directory).resolve()
    return source.parent


def review_species(path, project):
    """Resolve species from model provenance, with a condition-name fallback."""
    model = model_name(path)
    if project is not None:
        sys.path.insert(0, str(project))
        from carbon_flow_tools.smiles_completion import species_name
        for plan in (path.parent / "analysis_plan.json", project / "data/analysis_plan.json"):
            if plan.is_file():
                entry = json.loads(plan.read_text(encoding="utf-8")).get("models", {}).get(model)
                if entry is not None:
                    return species_name(model, entry.get("source_model"))
        return species_name(model)
    match = re.match(r"^([A-Z][a-z]+_[a-z][a-z0-9]+)(?:_|$)", model)
    return match.group(1) if match else model


def read_ai_records(directory, filename, fields):
    """CSV rows are the user-editable allowlist for reusing archived AI results."""
    source = directory / filename
    if not source.is_file():
        return []
    with source.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames != fields:
            raise ValueError("Unexpected AI record columns: " + str(source))
        rows = list(reader)
    if any(None in row or any(row[field] is None for field in fields) for row in rows):
        raise ValueError("Malformed AI CSV row: " + str(source))
    return [row for row in rows if any(value.strip() for value in row.values())]


def save_ai_records(directory, filename, fields, additions):
    """Merge new accepted rows with the current CSV, never rebuild from archives."""
    destination = directory / filename
    additions = read_ai_records(directory, filename, fields) + list(additions)
    rows = []
    for addition in additions:
        matching = [row for row in rows if all(row[field] == addition[field] for field in fields[:-1])]
        if matching:
            matching[0]["Source"] = " | ".join(dict.fromkeys(
                matching[0]["Source"].split(" | ") + addition["Source"].split(" | ")))
        else:
            rows.append(addition)
    rows.sort(key=lambda row: tuple(row[field] for field in fields[:-1]))
    directory.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".ai-records-", dir=str(directory)) as temporary:
        staged = Path(temporary) / filename
        with staged.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)
        staged.replace(destination)
    return destination


def index_ai_mappings(report, targets, proposals, directory, species):
    """Only successful AI replacements enter the index, in original equation direction."""
    accepted = {row["ReactionID"] for row in report["reactions"] if row["status"] == "accepted"}
    rows = []
    for proposal in proposals:
        key = identity(proposal["ReactionID"])
        if key not in accepted:
            continue
        mapped = proposal["MappedReactionSMILES"]
        if targets[key]["MappingDirection"] == "reverse":
            mapped = ">>".join(reversed(mapped.split(">>")))
        sources = [item["url"] for item in proposal["evidence"]]
        sources.append("archive: ai_mapping_reviews/" + Path(report["archive_file"]).name)
        rows.append(dict(Species=species, ReactionID=key, MappedReactionSMILES=mapped,
                         Source=" | ".join(dict.fromkeys(sources))))
    if rows:
        return save_ai_records(directory, "ai_atom_mappings.csv",
            ["Species", "ReactionID", "MappedReactionSMILES", "Source"], rows)


def archive_review(report, targets, proposals, directory):
    """Keep immutable review records; rejected proposals are not reusable mappings."""
    payload = dict(
        schema_version=1,
        input_file=report["input_file"],
        input_sha256=report["input_sha256"],
        output_file=report["output_file"],
        project_validation=report["project_validation"],
        reviews=report.get("reviews", []),
        reviewed_targets=[targets[identity(row["ReactionID"])]
                          for row in report.get("reviews", [])],
        reactions=[
            dict(
                target={
                    field: targets[key][field]
                    for field in (
                        "ReactionID",
                        "ReactionName",
                        "Equation",
                        "MappingDirection",
                        "ReactionSMILES",
                        "input_variants",
                        "partial_mappings",
                        "omitted_components",
                        "metabolites",
                    )
                },
                proposal=proposal,
                outcome=next(
                    row for row in report["reactions"] if row["ReactionID"] == key
                ),
            )
            for proposal in proposals
            for key in [identity(proposal["ReactionID"])]
            if key in targets
        ],
    )
    serialized = json.dumps(payload, indent=2, sort_keys=True, allow_nan=False) + "\n"
    key = hashlib.sha256(serialized.encode("utf-8")).hexdigest()
    directory.mkdir(parents=True, exist_ok=True)
    output = directory / (key + ".json")
    if output.exists():
        if output.read_text(encoding="utf-8") != serialized:
            raise ValueError(
                "Existing review archive differs; retained: " + str(output)
            )
        return str(output)
    with tempfile.TemporaryDirectory(
        prefix=".review-", dir=str(directory)
    ) as temporary:
        staged = Path(temporary) / output.name
        staged.write_text(serialized, encoding="utf-8")
        staged.replace(output)
    return str(output)


def cached_reviews(project, targets, species=None):
    """Expose exact-input candidates for review, never apply them automatically."""
    candidates, issues = {}, []
    directory = project / "data/ai_mapping_reviews" if project else None
    if directory is None or not directory.is_dir():
        return candidates, issues
    try:
        enabled = {(row["Species"], row["ReactionID"], row["MappedReactionSMILES"])
                   for row in read_ai_records(directory.parent, "ai_atom_mappings.csv",
                       ["Species", "ReactionID", "MappedReactionSMILES", "Source"])}
    except (OSError, ValueError) as error:
        return candidates, [dict(cache_file=str(directory.parent / "ai_atom_mappings.csv"), error=str(error))]
    fields = (
        "ReactionID",
        "ReactionName",
        "Equation",
        "MappingDirection",
        "input_variants",
    )
    for path in sorted(directory.glob("*.json")):
        try:
            archived = json.loads(path.read_text(encoding="utf-8"))
            if archived.get("schema_version") != 1:
                raise ValueError("Unsupported review archive schema")
            archived_species = review_species(Path(archived["input_file"]), project)
            if species is not None and archived_species != species:
                continue
            for record in archived["reactions"]:
                if record["outcome"]["status"] != "accepted":
                    continue
                target = record["target"]
                key = target["ReactionID"]
                mapped = record["proposal"]["MappedReactionSMILES"]
                if target["MappingDirection"] == "reverse":
                    mapped = ">>".join(reversed(mapped.split(">>")))
                # Keeping evidence does not undo a user's deletion or correction in the CSV.
                if (archived_species, key, mapped) not in enabled:
                    continue
                if (key in targets and target.get("partial_mappings", []) !=
                        targets[key].get("partial_mappings", [])):
                    continue
                if key not in targets or any(
                    target[field] != targets[key][field] for field in fields
                ):
                    continue
                candidates.setdefault(key, []).append(
                    dict(
                        archive_file=str(path),
                        input_sha256=archived["input_sha256"],
                        project_validation=archived["project_validation"],
                        proposal=record["proposal"],
                        checks=record["outcome"]["checks"],
                        requires_revalidation=True,
                    )
                )
        except (OSError, ValueError, KeyError, TypeError, AttributeError) as error:
            issues.append(dict(archive_file=str(path), error=str(error)))
    return candidates, issues


def apply_workbook(entry, project=None, mapping_skill=None, output_dir=None):
    source = Path(entry["input_file"]).resolve()
    model = model_name(source)
    before = digest(source)
    if entry.get("input_sha256") != before:
        raise ValueError("Input SHA256 changed since extraction")
    workbook = load_workbook(source)
    try:
        targets, issues = review_targets(workbook)
        proposals = entry.get("replacements", [])
        if not isinstance(proposals, list):
            raise ValueError("replacements must be a list")
        if any(not isinstance(row, dict) for row in proposals):
            raise ValueError("Each replacement must be an object")
        ids = [identity(row.get("ReactionID")) for row in proposals]
        if len(set(ids)) != len(ids):
            raise ValueError("Duplicated proposal ReactionID")
        reviews = entry.get("reviews", [])
        if not isinstance(reviews, list):
            raise ValueError("reviews must be a list")
        decisions = {}
        for row in reviews:
            if not isinstance(row, dict):
                raise ValueError("Each review must be an object")
            key = identity(row.get("ReactionID"))
            if key not in targets or key in decisions:
                raise ValueError("Unknown or duplicated review ReactionID: " + key)
            if row.get("decision") not in ("unchanged", "unresolved"):
                raise ValueError("Review decision must be unchanged or unresolved")
            if type(row.get("AIReviewed")) is not bool:
                raise ValueError("Review AIReviewed must be a Boolean")
            if (
                not isinstance(row.get("rationale"), str)
                or not row["rationale"].strip()
            ):
                raise ValueError("Each review needs a rationale")
            decisions[key] = row
        api = project_api(project)
        prepared = {}
        if api is not None:
            import pandas as pd

            with pd.ExcelFile(source, engine="openpyxl") as excel:
                names = api["input_sheets"](excel)
            reactions, _, metabolites = api["pipeline"].read_input(source, *names)
            checker = getattr(api["pipeline"], "carbon_structure_issues", None)
            if callable(checker):
                problems = checker(metabolites)
                for reaction in reactions:
                    conflicts = [c.metabolite for c in reaction.left + reaction.right
                                 if c.metabolite in problems]
                    if conflicts:
                        reaction.structure_error = "Input formula/structure conflict: " + "; ".join(
                            name + " (" + problems[name] + ")" for name in sorted(set(conflicts)))
            prepared = {reaction.reaction_id: reaction for reaction in reactions}
            for key, target in targets.items():
                reaction = prepared.get(key)
                if not target["input_variants"] and reaction is not None and reaction.reaction_smiles:
                    # Reconstruct only from the saved exact inputs, never from a guessed structure.
                    target["input_variants"] = [reaction.reaction_smiles]
                    target["ReactionSMILES"] = reaction.reaction_smiles
                    target["issues"] = [item for item in target["issues"] if item != "No usable unmapped ReactionSMILES"]
        validator = scientific_validator(mapping_skill)
        outcomes = {
            key: dict(
                ReactionID=key,
                status="unchanged",
                reason="No accepted proposal supplied",
            )
            for key in targets
        }
        accepted = 0
        for key, proposal in zip(ids, proposals):
            try:
                if key not in targets:
                    raise ValueError("Reaction is not a flagged review target")
                checks = check_proposal(
                    proposal, targets[key], validator, prepared, api
                )
                update_row(workbook, key, targets[key], proposal)
                outcomes[key] = dict(
                    ReactionID=key,
                    status="accepted",
                    MappedReactionSMILES=proposal["MappedReactionSMILES"],
                    checks=checks,
                    evidence=proposal["evidence"],
                    rationale=proposal["rationale"],
                    evidence_status=proposal["evidence_status"],
                    limitations=proposal.get("limitations", []),
                )
                accepted += 1
            except (ValueError, TypeError, KeyError) as error:
                outcomes[key] = dict(
                    ReactionID=key, status="rejected", reason=str(error)
                )
        output = None
        reviewed = {key for key, row in outcomes.items() if row["status"] == "accepted"}
        reviewed.update(key for key, row in decisions.items() if row["AIReviewed"])
        for key, row in decisions.items():
            if outcomes[key]["status"] == "accepted":
                raise ValueError(
                    "Accepted reaction also has a non-replacement review: " + key
                )
            if (
                outcomes[key]["status"] == "rejected"
                and row["decision"] != "unresolved"
            ):
                raise ValueError("A rejected mapping must remain unresolved: " + key)
        for key, decision in decisions.items():
            apply_rule_verdict(workbook, key, targets[key], decision)
        if any(row.get("rule_verdict") == "accepted" for row in decisions.values()):
            if api is None:
                raise ValueError("Accepting saved partial rules requires project validation")
            api["pipeline"].load_final_mappings(source, {})
        if accepted or decisions:
            mark_reviewed(workbook, reviewed)
            directory = output_directory(output_dir, project, source)
            output = directory / (model + "_2_atom_mapping_AIGenerated.xlsx")
            if output.resolve() == source or output.exists():
                raise ValueError("Existing input/output retained: " + str(output))
            directory.mkdir(parents=True, exist_ok=True)
            with tempfile.TemporaryDirectory(
                prefix=".carbon-review-", dir=str(directory)
            ) as temporary:
                staged = Path(temporary) / output.name
                if api is not None:
                    api["pipeline"].save_output(workbook, staged)
                else:
                    workbook.save(staged)
                checked = load_workbook(staged)
                try:
                    if not same_sheet_data(workbook, checked):
                        raise ValueError(
                            "Saved workbook differs from intended sheet data"
                        )
                finally:
                    checked.close()
                if digest(source) != before:
                    raise ValueError("Input changed during review; no output committed")
                if output.exists():
                    raise ValueError(
                        "An output appeared during review; existing file retained"
                    )
                staged.replace(output)
        report = dict(
            input_file=str(source),
            input_sha256=before,
            output_file=str(output) if output else None,
            accepted=accepted,
            flagged=len(targets),  # Legacy field: target union, not TRUE flags alone.
            target_counts=target_counts(targets),
            ai_reviewed=len(reviewed),
            reviews=reviews,
            project_validation=("passed" if accepted or any(
                row.get("rule_verdict") == "accepted" for row in decisions.values()) else "not_run")
            if api
            else "not_available",
            issues=issues,
            reactions=[
                dict(
                    row,
                    flagged=key in targets,
                    target={
                        field: targets[key][field]
                        for field in (
                            "ReactionName",
                            "Equation",
                            "MappingDirection",
                            "ReactionSMILES",
                            "partial_mappings",
                            "metabolites",
                            "final",
                            "review_trigger",
                            "missing_final_row",
                        )
                    }
                    if key in targets
                    else None,
                )
                for key, row in outcomes.items()
            ],
        )
        report.update(review_progress(report))
        try:
            report["structure_changes"] = completed_structure_changes(source, before)
        except (OSError, ValueError, KeyError, TypeError) as error:
            report["structure_changes"] = []
            report["structure_report_error"] = str(error)
        if proposals or reviews:
            directory = (
                project if project else output_directory(output_dir, project, source)
            ) / "data/ai_mapping_reviews"
            try:
                report["archive_file"] = archive_review(
                    report, targets, proposals, directory
                )
                index_ai_mappings(report, targets, proposals, directory.parent,
                                  review_species(source, project))
            except (OSError, ValueError, TypeError) as error:
                report["archive_error"] = str(error)
        return report
    finally:
        workbook.close()



def completed_structure_changes(source, workbook_hash):
    """Read only structure additions bound to this exact staging workbook and archive."""
    handoff = source.with_suffix(".structures.json")
    if not handoff.is_file():
        return []
    metadata = json.loads(handoff.read_text(encoding="utf-8"))
    if metadata["output_sha256"] != workbook_hash:
        raise ValueError("Structure handoff does not match the reviewed workbook")
    archive = Path(metadata["structure_archive"])
    if digest(archive) != archive.stem:
        raise ValueError("Structure archive hash mismatch")
    data = json.loads(archive.read_text(encoding="utf-8"))
    if data["parent_workbook_sha256"] != metadata["input_sha256"]:
        raise ValueError("Structure archive does not match the handoff parent")
    changes = [dict(MetaboliteID=row["identity"]["MetaboliteID"],
                    MetaboliteName=row["identity"]["MetaboliteName"],
                    SMILES=row["proposal"]["SMILES"], action="Added")
               for row in data["structures"]]
    if len(changes) != metadata["completed_structures"]:
        raise ValueError("Structure change count differs from the verified handoff")
    return changes


def write_delivery_report(reports, directory):
    """Keep the readable report to a summary and applied changes; JSON keeps the audit."""
    directory.mkdir(parents=True, exist_ok=True)
    stem = "AI_mapping_review_" + datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    json_path, text_path = directory / (stem + ".json"), directory / (stem + ".md")
    fields = ("scientifically_reviewed", "replaced", "retained", "unresolved")
    lines = ["# AI-assisted atom mapping review", "",
        "| Model | AI reviewed | Replaced | Retained | Reviewed unresolved | Status |",
        "| --- | " + " | ".join(["---:"] * len(fields)) + " | --- |"]
    totals = {key: 0 for key in fields}
    structures, mappings, outputs, notices = [], [], [], []
    incomplete = False
    def cell(value):
        return str(value or "").replace("|", "\\|").replace("\n", " ").replace("\r", " ")
    for report in reports:
        name = Path(report.get("input_file") or "Unknown input").name
        if report.get("error"):
            incomplete = True
            lines.append("| " + cell(name) + " | " + " | ".join(["unknown"] * len(fields)) + " | failed |")
            notices.append(name + ": input processing failed; see JSON.")
            continue
        counts = report["review_counts"]
        incomplete |= report["review_status"] == "incomplete"
        for key in fields:
            totals[key] += counts[key]
        lines.append("| {} | {} | {} |".format(cell(name),
            " | ".join(str(counts[key]) for key in fields), report["review_status"]))
        if report["remaining_review_ids"]:
            notices.append(name + ": review is not complete; {} reactions still require analysis.".format(
                len(report["remaining_review_ids"])))
        if any(report.get(key) for key in ("issues", "archive_error", "structure_report_error")):
            notices.append(name + ": delivery or validation warnings; see JSON.")
        for change in report.get("structure_changes", []):
            structures.append("| " + " | ".join(cell(value) for value in
                (name, change["MetaboliteID"], change["MetaboliteName"],
                 change["action"], change["SMILES"])) + " |")
        for outcome in report["reactions"]:
            if outcome["status"] != "accepted" or not outcome["flagged"]:
                continue
            target = outcome.get("target") or {}
            mappings.append("| " + " | ".join(cell(value) for value in
                (name, outcome["ReactionID"], target.get("ReactionName", ""),
                 target.get("Equation", ""))) + " |")
        if report.get("output_file"):
            outputs.append("- " + str(report["output_file"]))
    status = ("incomplete" if incomplete else "complete"
              if any(report.get("target_counts", {}).get("targets") for report in reports)
              else "no_targets")
    lines += ["| Total | " + " | ".join(str(totals[key]) for key in fields) + " | " + status + " |", ""]
    lines += [notice for notice in notices] + ([""] if notices else [])
    lines += ["## SMILES changes", ""]
    if structures:
        lines += ["| Model | Metabolite ID | Metabolite | Change | SMILES |",
                  "| --- | --- | --- | --- | --- |"] + structures + [""]
    else:
        lines += ["No structure changes were recorded for this review.", ""]
    lines += ["## Applied mapping changes", ""]
    if mappings:
        lines += ["| Model | ReactionID | Reaction | Equation |",
                  "| --- | --- | --- | --- |"] + mappings + [""]
    else:
        lines += ["No mapping replacements were applied.", ""]
    lines += ["## Files", ""] + outputs + ["",
              "Full mapped SMILES are in the reviewed workbooks. Validation, evidence and "
              "continuation records: [JSON](" + json_path.name + ").", ""]
    with json_path.open("x", encoding="utf-8") as stream:
        json.dump({"review_status": status, "totals": totals, "workbooks": reports},
                  stream, indent=2, allow_nan=False)
    with text_path.open("x", encoding="utf-8") as stream:
        stream.write("\n".join(lines))
    return dict(markdown=str(text_path), json=str(json_path))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command")
    inspect = commands.add_parser("inspect", help="Extract flagged reaction inputs")
    inspect.add_argument("workbooks", nargs="+")
    apply = commands.add_parser(
        "apply", help="Validate proposals and write separate copies"
    )
    apply.add_argument("proposals", help="JSON with a workbooks list")
    for command in (inspect, apply):
        command.add_argument("--project", help="Compatible local project root")
    apply.add_argument(
        "--mapping-skill", help="Installed enzymatic-atom-mapping directory"
    )
    apply.add_argument(
        "--output-dir", help="Delivery directory; existing files are never overwritten"
    )
    args = parser.parse_args(argv)
    if args.command is None:
        parser.error("Choose inspect or apply")
    if args.command == "inspect":
        entries = [{"input_file": str(Path(p).resolve())} for p in args.workbooks]
    else:
        entries = json.loads(Path(args.proposals).read_text(encoding="utf-8"))[
            "workbooks"
        ]
    if not isinstance(entries, list) or not entries:
        raise ValueError("Supply at least one workbook")
    report = []
    for entry in entries:
        if not isinstance(entry, dict):
            report.append(dict(error="Each workbook entry must be an object"))
            continue
        try:
            source = Path(entry["input_file"]).resolve()
            project = find_project(args.project, [source])
            if args.command == "inspect":
                model = model_name(source)
                before = digest(source)
                workbook = load_workbook(source)
                try:
                    targets, issues = review_targets(workbook)
                finally:
                    workbook.close()
                if digest(source) != before:
                    raise ValueError("Input changed during extraction; retry")
                cached, cache_issues = cached_reviews(project, targets, review_species(source, project))
                for key, target in targets.items():
                    target["cached_candidates"] = cached.get(key, [])
                report.append(
                    dict(
                        input_file=str(source),
                        input_sha256=before,
                        model=model,
                        targets=list(targets.values()),
                        target_counts=target_counts(targets),
                        issues=issues,
                        cache_issues=cache_issues,
                        project=str(project) if project else None,
                        step3_available=bool(
                            project and (project / "step3_carbon_flow.py").is_file()
                        ),
                        step4_available=bool(
                            project and (project / "step4_carbon_flow_ai.py").is_file()
                        ),
                    )
                )
            else:
                report.append(
                    apply_workbook(entry, project, args.mapping_skill, args.output_dir)
                )
        except (
            ValueError,
            TypeError,
            KeyError,
            OSError,
            ImportError,
            BadZipFile,
        ) as error:
            report.append(dict(input_file=entry.get("input_file"), error=str(error)))
    delivery = {"workbooks": report}
    if args.command == "apply":
        delivery["review_status"] = (
            "incomplete" if any(row.get("error") or row.get("review_status") == "incomplete"
                                for row in report)
            else "complete" if any(row["target_counts"]["targets"] for row in report)
            else "no_targets")
        try:
            source = Path(entries[0]["input_file"]).resolve()
            project = find_project(args.project, [source])
            delivery["report_files"] = write_delivery_report(
                report, output_directory(args.output_dir, project, source)
            )
        except (OSError, ValueError, KeyError, TypeError) as error:
            delivery["report_error"] = str(error)
    print(json.dumps(delivery, indent=2, default=str, allow_nan=False))
    return (
        2
        if delivery.get("report_error")
        or any(
            row.get("error")
            or row.get("review_status") == "incomplete"
            or row.get("archive_error")
            or row.get("structure_report_error")
            or row.get("issues")
            or any(r["status"] == "rejected" for r in row.get("reactions", []))
            for row in report
        )
        else 0
    )


if __name__ == "__main__":
    try:
        exit_code = main()
    except (ValueError, TypeError, KeyError, OSError, ImportError) as error:
        print(json.dumps({"error": str(error)}, indent=2))
        exit_code = 2
    sys.exit(exit_code)
