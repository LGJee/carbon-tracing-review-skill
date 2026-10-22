"""Synthetic structure-completion tests; sources and mapper results are fixtures."""

from copy import deepcopy
import importlib.util
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

import complete_structures as completion
import review_workbooks as review

PROJECT = Path(sys.argv[1]).resolve()
MAPPING_SKILL = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else None
sys.path.insert(0, str(PROJECT))
review.project_api(PROJECT)
from carbon_flow_tools.pipeline import finish_analysis, load_final_mappings
from carbon_flow_tools.reactions import prepare_reactions

MAPPED = "[CH3:1][CH2:2][OH:3]>>[CH3:1][CH2:2][OH:3]"


def proposal(key="b", name="B[c]", smiles="CCO"):
    return dict(MetaboliteID=key, MetaboliteName=name, SMILES=smiles,
                database_id="Synthetic fixture, not scientific evidence", rationale="Synthetic identity test",
                evidence=[dict(url="https://example.org/test-fixture", accessed_on="2026-10-18",
                               supported_claim="Synthetic test only")])


def fixture(root, product="B[c]", reaction_equation=None):
    source = pd.DataFrame([
        ["T", "Transport", 1, "A[e] => A[c]"],
        ["R", "Newly mappable", 1, "A[c] => B[c]"],
        ["U", "Unchanged", 1, "C[c] => A[c]"],
        ["P", "Still blocked", 1, "B[c] => protein[c]"],
        ["E", "Boundary", 1, "B[c] => "]],
        columns=["ReactionID", "ReactionName", "Flux", "Equation"])
    metabolites = pd.DataFrame([
        ["a_e", "A[e]", "CCO", "C2H6O", 0], ["a", "A[c]", "CCO", "C2H6O", 0],
        ["b", "B[c]", "", "C2H6O", 0], ["c", "C[c]", "CCO", "C2H6O", 0],
        ["pool", "protein[c]", "", "", ""]],
        columns=["MetaboliteID", "MetaboliteName", "SMILES_Model", "MolecularFormula", "Charge"])
    if reaction_equation is not None:
        source.loc[source.ReactionID.eq("R"), "Equation"] = reaction_equation
    source["Equation"] = source.Equation.str.replace("B[c]", product, regex=False)
    metabolites.loc[metabolites.MetaboliteID.eq("b"), "MetaboliteName"] = product
    path = root / "input.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        source.to_excel(writer, sheet_name="Nonzero", index=False)
        metabolites.to_excel(writer, sheet_name="Metabolites", index=False)
    predictions = {backend: prepare_reactions(source, metabolites) for backend in review.BACKENDS}
    for rows in predictions.values():
        rows[2].mapped_smiles, rows[2].status = MAPPED, "mapped"
    finish_analysis("Test", dict(input_file=path, metabolite_sheet="Metabolites"),
                    metabolites, predictions, [dict(Mapper="Input", SourceInputSHA256=review.digest(path))],
                    1, list(predictions), root, stage="map")
    return root / "Test_2_atom_mapping.xlsx"


class CompletionTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        index_writer = review.index_ai_mappings
        def save_index(report, targets, proposals, output_dir, species):
            if output_dir == PROJECT / "data":
                output_dir = Path(directory.name) / "data"
            return index_writer(report, targets, proposals, output_dir, species)
        patcher = patch.object(review, "index_ai_mappings", side_effect=save_index)
        patcher.start()
        self.addCleanup(patcher.stop)
    def test_completed_reduced_input_is_remapped_and_still_requires_review(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = fixture(root, reaction_equation="A[c] + protein[c] => B[c] + protein[c]")
            before = review.digest(path)
            project = root / "project"
            project.mkdir()
            shutil.copyfile(PROJECT / "step2_atom_mapping.py", project / "step2_atom_mapping.py")
            original_context = completion.context
            calls = []
            def predict(model, reactions, backend, settings, threads):
                calls.append((backend, [r.reaction_id for r in reactions if r.status == "ready"]))
                for reaction in reactions:
                    if reaction.status == "ready":
                        reaction.mapped_smiles, reaction.status = MAPPED, "mapped"
            with patch.object(completion, "context", side_effect=lambda path, _: original_context(path, PROJECT)), \
                 patch("carbon_flow_tools.mapper_runner.validate_mapper_config"), \
                 patch("carbon_flow_tools.mapper_runner.run_mapper", side_effect=predict):
                result = completion.complete(dict(input_file=str(path), input_sha256=before,
                    structures=[proposal()]), project, root / "staging")
            self.assertEqual(len(calls), 3)
            self.assertTrue(all("R" in ids for _, ids in calls))
            workbook = load_workbook(result["output_file"])
            try:
                targets, _ = review.review_targets(workbook)
                self.assertIn("R", targets)
                self.assertEqual(targets["R"]["final"]["MappingCoverage"], "Partial")
                self.assertTrue(targets["R"]["final"]["ReviewRequired"])
                self.assertFalse(targets["R"]["final"]["AIGenerated"])
                self.assertEqual(len(targets["R"]["omitted_components"]), 2)
            finally:
                workbook.close()
            self.assertEqual(review.digest(path), before)

    def test_identity_validation_rejects_guesses_and_existing_structures(self):
        row = SimpleNamespace(MetaboliteID="b", MetaboliteName="B[c]", MolecularFormula="C2H6O", Charge=0)
        self.assertEqual(completion.validate_structure(row, proposal()), "CCO")
        for changed in (dict(SMILES="CC"), dict(SMILES="*"), dict(evidence=[]),
                        dict(MetaboliteName="Other[c]"), dict(SMILES="[NH4+]")):
            with self.assertRaises(ValueError):
                completion.validate_structure(row, dict(proposal(), **changed))
        row.SMILES_Model = "CCO"
        with self.assertRaises(ValueError):
            completion.validate_structure(row, proposal())
        row = SimpleNamespace(MetaboliteID="pool", MetaboliteName="protein[c]")
        with self.assertRaises(ValueError):
            completion.validate_structure(row, proposal("pool", "protein[c]"))

    def test_cli_inspection_reports_missing_identities_without_mapper_calls(self):
        with tempfile.TemporaryDirectory() as directory:
            path = fixture(Path(directory))
            text = subprocess.check_output([sys.executable,
                   str(Path(completion.__file__)), "inspect", str(path), "--project", str(PROJECT)],
                   universal_newlines=True)
            report = json.loads(text)
            self.assertEqual(report["input_sha256"], review.digest(path))
            self.assertEqual({row["MetaboliteID"] for row in report["missing_structures"]}, {"b", "pool"})

    def test_cache_is_durable_idempotent_and_conflicts_are_not_overwritten(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            identity = completion.identity_context("Test", SimpleNamespace(MetaboliteID="b", MetaboliteName="B[c]"))
            accepted = [dict(identity=identity, proposal=proposal())]
            archive = completion.persist(root, "Test", accepted, "0" * 64)
            self.assertEqual(archive, completion.persist(root, "Test", accepted, "0" * 64))
            self.assertEqual(len(list((root / "data/ai_structure_reviews").glob("*.json"))), 1)
            before = (root / "data/ai_smiles.csv").read_bytes()
            bad = [dict(identity=identity, proposal=proposal(smiles="CO"))]
            with self.assertRaises(ValueError):
                completion.persist(root, "Test", bad, "0" * 64)
            self.assertEqual((root / "data/ai_smiles.csv").read_bytes(), before)
            self.assertEqual(len(before.decode().splitlines()), 2)

    def test_species_ai_table_is_shared_by_conditions_and_compartments(self):
        import csv
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for model, key, name in (("Escherichia_coli_non_IUPD", "first", "B[c]"),
                                      ("Escherichia_coli_IUPD", "different", "B[e]")):
                identity = completion.identity_context(model, SimpleNamespace(MetaboliteID=key, MetaboliteName=name))
                completion.persist(root, model, [dict(identity=identity, proposal=proposal(key, name))], "0" * 64)
            with (root / "data/ai_smiles.csv").open(encoding="utf-8-sig") as handle:
                reader = csv.DictReader(handle)
                self.assertEqual(reader.fieldnames, ["Species", "MetaboliteName", "SMILES", "Source"])
                rows = list(reader)
            self.assertEqual(len(rows), 1)
            self.assertEqual({key: value for key, value in rows[0].items() if key != "Source"},
                             dict(Species="Escherichia_coli", MetaboliteName="B", SMILES="CCO"))
            self.assertIn("https://example.org/test-fixture", rows[0]["Source"])
            self.assertFalse((root / "data/manual_smiles.csv").exists())

    def test_manual_structure_table_is_never_changed_by_ai(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            manual = root / "data/manual_smiles.csv"
            manual.parent.mkdir()
            manual.write_text("Species,MetaboliteName,SMILES\nTest,Other,CO\n")
            before = manual.read_bytes()
            identity = completion.identity_context("Test", SimpleNamespace(MetaboliteID="b", MetaboliteName="B[c]"))
            completion.persist(root, "Test", [dict(identity=identity, proposal=proposal())], "0" * 64)
            self.assertEqual(manual.read_bytes(), before)
            self.assertTrue((root / "data/ai_smiles.csv").is_file())

    def test_missing_mapper_environment_keeps_discoveries_but_no_delivery(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = fixture(root)
            original_hash = review.digest(path)
            project = root / "project"
            project.mkdir()
            shutil.copyfile(PROJECT / "step2_atom_mapping.py", project / "step2_atom_mapping.py")
            entry = dict(input_file=str(path), input_sha256=original_hash, structures=[proposal()])
            original_context = completion.context
            with patch.object(completion, "context", side_effect=lambda path, _: original_context(path, PROJECT)), \
                 patch("carbon_flow_tools.mapper_runner.validate_mapper_config", side_effect=FileNotFoundError("Missing test environment")):
                with self.assertRaises(FileNotFoundError):
                    completion.complete(entry, project, root / "staging")
            self.assertEqual(review.digest(path), original_hash)
            self.assertFalse((root / "staging" / path.name).exists())
            self.assertTrue((project / "data/ai_smiles.csv").is_file())
            self.assertEqual(len(list((project / "data/ai_structure_reviews").glob("*.json"))), 1)

    def test_empty_mapper_results_still_allow_complete_ai_replacement(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = fixture(root)
            project = root / "project"
            project.mkdir()
            shutil.copyfile(PROJECT / "step2_atom_mapping.py", project / "step2_atom_mapping.py")
            entry = dict(input_file=str(path), input_sha256=review.digest(path), structures=[proposal()])
            original_context = completion.context
            def empty(model, reactions, backend, settings, threads):
                for reaction in reactions:
                    if reaction.status == "ready":
                        reaction.mapped_smiles = ""
                        reaction.status = backend + " failed: Empty result"
            with patch.object(completion, "context", side_effect=lambda path, _: original_context(path, PROJECT)), \
                 patch("carbon_flow_tools.mapper_runner.validate_mapper_config"), \
                 patch("carbon_flow_tools.mapper_runner.run_mapper", side_effect=empty):
                result = completion.complete(entry, project, root / "staging")
            staged = Path(result["output_file"])
            workbook = load_workbook(staged)
            try:
                targets, _ = review.review_targets(workbook)
                self.assertIn("R", targets)
                self.assertTrue(targets["R"]["ReactionSMILES"])
                self.assertTrue(all(not row["MappedReactionSMILES"] for row in targets["R"]["predictions"]))
                self.assertFalse(targets["R"]["issues"])
            finally:
                workbook.close()
            replacement = dict(ReactionID="R", MappedReactionSMILES=MAPPED, scope="all",
                evidence_status="evidence-supported", rationale="Synthetic fixture; not real scientific evidence",
                evidence=[dict(url="https://example.org/test-fixture", accessed_on="2026-10-18",
                               claim="Synthetic correspondence for handoff validation only")])
            ai_entry = dict(input_file=str(staged), input_sha256=review.digest(staged), replacements=[replacement])
            with patch.object(review, "archive_review", return_value="synthetic-test-only"):
                delivered = review.apply_workbook(ai_entry, PROJECT, mapping_skill=MAPPING_SKILL, output_dir=root / "delivery")
            self.assertEqual(delivered["accepted"], 1, delivered["reactions"])
            output = Path(delivered["output_file"])
            workbook = load_workbook(output)
            try:
                _, final = review.table(workbook, "Final_Mapping", review.FINAL_COLUMNS)
                self.assertEqual(final["R"]["MappedReactionSMILES"], MAPPED)
                self.assertTrue(final["R"]["AIGenerated"])
                self.assertTrue(final["R"]["AIReviewed"])
                for backend in review.BACKENDS:
                    _, raw = review.table(workbook, backend, {"ReactionID", "MappedReactionSMILES"})
                    self.assertFalse(raw["R"]["MappedReactionSMILES"])
            finally:
                workbook.close()
            loaded, _, _ = load_final_mappings(output, {})
            self.assertEqual(next(row for row in loaded if row.reaction_id == "R").status, "mapped")

    def test_identity_transport_needs_no_mapper_call(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = fixture(root, product="A[m]")
            project = root / "project"
            project.mkdir()
            shutil.copyfile(PROJECT / "step2_atom_mapping.py", project / "step2_atom_mapping.py")
            entry = dict(input_file=str(path), input_sha256=review.digest(path), structures=[proposal(name="A[m]")])
            original_context = completion.context
            with patch.object(completion, "context", side_effect=lambda path, _: original_context(path, PROJECT)), \
                 patch("carbon_flow_tools.mapper_runner.run_mapper", side_effect=AssertionError("Transport needs no mapper")):
                result = completion.complete(entry, project, root / "staging")
            self.assertEqual(result["mapper_reactions"], [])
            loaded, _, _ = load_final_mappings(Path(result["output_file"]), {})
            self.assertTrue(next(r for r in loaded if r.reaction_id == "R").is_transport)

    def test_incremental_remapping_preserves_other_rows_and_step3_contract(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = fixture(root)
            original_hash = review.digest(path)
            project = root / "project"
            project.mkdir()
            shutil.copyfile(PROJECT / "step2_atom_mapping.py", project / "step2_atom_mapping.py")
            entry = dict(input_file=str(path), input_sha256=original_hash, structures=[proposal()])
            with self.assertRaises(ValueError):
                completion.complete(dict(entry, input_sha256="0" * 64), PROJECT, root / "bad")
            original_context = completion.context
            calls = []
            def mapped(model, reactions, backend, settings, threads):
                calls.append((backend, [r.reaction_id for r in reactions if r.status == "ready"]))
                for reaction in reactions:
                    if reaction.status == "ready":
                        reaction.mapped_smiles, reaction.status = MAPPED, "mapped"
            with patch.object(completion, "context", side_effect=lambda path, _: original_context(path, PROJECT)), \
                 patch("carbon_flow_tools.mapper_runner.validate_mapper_config"), \
                 patch("carbon_flow_tools.mapper_runner.run_mapper", side_effect=mapped):
                result = completion.complete(entry, project, root / "staging")
                cached = completion.inspect(path, project)
            self.assertEqual({name for name, _ in calls}, set(review.BACKENDS))
            self.assertTrue(all(ids == ["R"] for _, ids in calls))
            self.assertEqual(result["mapper_reactions"], ["R"])
            self.assertEqual(review.digest(path), original_hash)
            self.assertEqual(cached["missing_structures"][0]["cached_candidates"][0]["proposal"]["SMILES"], "CCO")
            index = project / "data/ai_smiles.csv"
            previous = index.read_bytes()
            index.write_text("Species,MetaboliteName,SMILES,Source\n", encoding="utf-8-sig")
            with patch.object(completion, "context", side_effect=lambda path, _: original_context(path, PROJECT)):
                without_cache = completion.inspect(path, project)
            self.assertTrue(all(not row["cached_candidates"] for row in without_cache["missing_structures"]))
            self.assertTrue(list((project / "data/ai_structure_reviews").glob("*.json")))
            index.write_bytes(previous)
            output = Path(result["output_file"])
            loaded, _, _ = load_final_mappings(output, {})
            self.assertEqual(next(r for r in loaded if r.reaction_id == "R").status, "mapped")
            review_entry = dict(input_file=str(output), input_sha256=review.digest(output), replacements=[],
                                reviews=[dict(ReactionID="R", decision="unchanged", AIReviewed=False,
                                              rationale="Automatic test triage only, not scientific review")])
            with patch.object(review, "archive_review", return_value="synthetic-test-only"):
                reviewed = review.apply_workbook(review_entry, PROJECT, mapping_skill=MAPPING_SKILL,
                                             output_dir=root / "delivery")
            self.assertTrue(Path(reviewed["output_file"]).name.endswith("_AIGenerated.xlsx"))
            load_final_mappings(Path(reviewed["output_file"]), {})
            original, updated = load_workbook(path), load_workbook(output)
            try:
                self.assertEqual(list(original["Nonzero"].values), list(updated["Nonzero"].values))
                old_mets, new_mets = list(original["Metabolites"].values), list(updated["Metabolites"].values)
                for old_row in old_mets[1:]:
                    if old_row[0] != "b":
                        new_row = next(row for row in new_mets[1:] if row[0] == old_row[0])
                        self.assertEqual(old_row, new_row[:len(old_mets[0])])
                for name in (*review.BACKENDS, "Mapping_Selection", "Final_Mapping"):
                    old = list(original[name].values)
                    new = list(updated[name].values)
                    for key in ("T", "U"):
                        self.assertEqual(next(row for row in old[1:] if row[0] == key),
                                         next(row for row in new[1:] if row[0] == key))
                _, final = review.table(updated, "Final_Mapping", review.FINAL_COLUMNS)
                self.assertFalse(final["R"]["AIGenerated"])
                self.assertTrue(final["R"]["ReviewRequired"])
                self.assertEqual(final["R"]["MappedReactionSMILES"], MAPPED)
                self.assertEqual(final["P"]["MappingCoverage"], "Unresolved")
                self.assertEqual(updated["Mapping_Info"].cell(2, 2).value, original["Mapping_Info"].cell(2, 2).value)
            finally:
                original.close()
                updated.close()
            with self.assertRaises(ValueError):
                completion.complete(entry, project, output.parent)


if __name__ == "__main__":
    unittest.main(argv=[sys.argv[0]])